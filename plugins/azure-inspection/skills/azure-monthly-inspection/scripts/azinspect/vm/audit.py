"""Independent re-check of the VM report: every number in the workbook and Markdown is recomputed from raw data."""
from __future__ import annotations

import json
import math
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal

from markdown_it import MarkdownIt
from openpyxl import load_workbook

from ..core.excel import unrecalculated_formula
from ..core.io import read_json, write_json
from .report import STEM

PAIRS = [('VM Uncached IOPS Consumed Percentage', 'UncachedDiskIOPS'), ('VM Uncached Bandwidth Consumed Percentage', 'UncachedDiskBytesPerSecond'),
         ('VM Cached IOPS Consumed Percentage', 'CombinedTempDiskAndCachedIOPS'), ('VM Cached Bandwidth Consumed Percentage', 'CombinedTempDiskAndCachedReadBytesPerSecond')]


def same(a, b):
    assert math.isclose(a, b, abs_tol=1e-6, rel_tol=1e-9), (a, b)


def shown(n):
    # same rounding as Excel's TEXT() (15 significant digits, then half up), see table.money
    return format(Decimal(format(n, '.15g')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP), ',.2f')


def audit(ws, log=print):
    f = ws.month_dir('findings') / 'vm'
    vs, rs, cs, ps = (read_json(f / n) for n in ('vms-analyzed.json', 'recommendations.json', 'candidate-checks.json', 'pools.json'))
    subs, manifest, rows = read_json(f / 'subscriptions.json'), read_json(f / 'collection-manifest.json'), read_json(f / 'rows.json')
    ids = {s['id'] for s in subs}
    assert len(ids) == len(subs), 'duplicate subscriptions'
    assert len(vs) == len(rs) == len({v['id'].lower() for v in vs}), 'VM and recommendation counts differ'
    assert all(v['subscriptionId'] in ids for v in vs), 'VM outside the subscription list'
    assert manifest['interval'] == 'P1D' and manifest['aggregation'] == 'Average,Count'
    vlookup, rlookup = {v['id'].lower(): v for v in vs}, {r['vmId'].lower(): r for r in rs}
    for v in vs:
        for _name, m in v['metrics'].items():
            valid = [d for d in m.get('daily', []) if d.get('average') is not None and (d.get('count') or 0) > 0]
            count = sum(d['count'] for d in valid)
            assert m['sampleCount'] == count and m['validDays'] == len(valid)
            expected = sum(d['average'] * d['count'] for d in valid) / count if count else None
            if expected is None:
                assert m['average'] is None, 'missing metric filled with a value'
            else:
                same(m['average'], expected)
        if v['cpuAverage'] is not None:
            same(v['cpuAverage'], v['metrics']['Percentage CPU']['average'])
        if v['memoryAverage'] is not None:
            same(v['memoryAverage'], 100 - v['metrics']['Available Memory Percentage']['average'])
    for c in cs:
        r, v = rlookup[c['vmId'].lower()], vlookup[c['vmId'].lower()]
        if c['passed']:
            cpu, memory = c['cores'] / r['currentCores'], c['memoryGiB'] / r['currentMemoryGiB']
            assert cpu in (.5, 1) and memory in (.5, 1) and (cpu, memory) != (1, 1)
            if cpu < 1:
                assert r['cpuAverage'] < 20
            if memory < 1:
                assert r['memoryAverage'] < 20
            same(c['projectedCPUPercent'], r['cpuAverage'] / cpu)
            same(c['projectedMemoryPercent'], r['memoryAverage'] / memory)
            assert c['priceUSDHour'] < r['currentPriceUSDHour'] and not c['checks']
            for metric, key in PAIRS:
                predicted = c['projectedIO'].get(metric)
                if predicted is None:
                    assert c['ioApplicability'][metric].startswith('不适用')
                    continue
                assert predicted <= 50 + 1e-9 and v['metrics'][metric]['sufficient']
                same(predicted, v['metrics'][metric]['average'] * float(r['currentCaps'][key]) / float(c['caps'][key]))
            assert len(v['dataDisks']) <= int(c['caps']['MaxDataDiskCount'])
    for r in rs:
        if r['status'] == '建议试点调整':
            assert r['state'] == 'PowerState/running' and r['candidate']['passed'] and not r['candidate']['checks']
            same(r['monthlySavingsUSD'], (r['currentPriceUSDHour'] - r['candidate']['priceUSDHour']) * 730)
            assert r['candidate']['sku'] == r['targetSKU']
        else:
            assert r['monthlySavingsUSD'] == 0
    for p in ps:
        members = [rlookup[i.lower()] for i in p['members'] if rlookup[i.lower()]['state'] == 'PowerState/running']
        assert p['runningCount'] == len(members)
        assert len({vlookup[r['vmId'].lower()]['subscriptionId'] for r in members}) <= 1
        same(p['monthlySavingsUSD'], sum(r['monthlySavingsUSD'] for r in members))
        if p['status'] == '建议整池试点':
            assert all(r['status'] == '建议试点调整' and r['targetSKU'] == p['targetSKU'] for r in members)
        else:
            assert not any(r['status'] == '建议试点调整' for r in members)
    assert len(rows) == len(vs)
    retail = [item for p in ws.raw('vm/prices').glob('*.json') for item in json.loads(p.read_text(encoding='utf-8')).get('items', [])]

    def validate_quote(q, key):
        if q is None:
            return
        matches = [p for p in retail if p['meterId'] == q['meterId'] and p['armRegionName'] == q['region'] and p['armSkuName'].lower() == q['armSkuName'].lower()]
        assert matches, q
        if key.startswith('ri'):
            term = '1 Year' if key == 'ri1' else '3 Years'
            assert any(p['type'] == 'Reservation' and p.get('reservationTerm') == term and math.isclose(p['retailPrice'], q['retailPrice'], abs_tol=1e-8) for p in matches)
        elif key.startswith('sp'):
            term = '1 Year' if key == 'sp1' else '3 Years'
            assert any(p['type'] == 'Consumption' and any(sp['term'] == term and math.isclose(sp['retailPrice'], q['retailPrice'], abs_tol=1e-8) for sp in p.get('savingsPlan', [])) for p in matches)
        else:
            assert any(p['type'] == 'Consumption' and math.isclose(p['retailPrice'], q['retailPrice'], abs_tol=1e-8) for p in matches)

    reports = ws.month_dir('reports')
    verification = read_json(f / 'verification.json')
    assert verification['status'] == 'success' and verification['totalErrors'] == 0
    recalculated = verification['recalculatedInExcel']
    w = load_workbook(reports / f'{STEM}.xlsx', data_only=recalculated)
    sheet, calc = w['VM建议指标价格总表'], w['价格计算依据']
    assert sheet.max_row == len(rows) + 4 and sheet.cell(4, 1).value == '订阅' and '公开列表价' in sheet['A2'].value
    total = sum(r['monthlySavingsUSD'] for r in rs)
    if recalculated:
        same(w['总览']['B12'].value, total)
        same(w['总览']['B13'].value, total * 12)
    for i, u in enumerate(rows, 5):
        r, v = rlookup[u['vmId'].lower()], vlookup[u['vmId'].lower()]
        assert sheet.cell(i, 1).value == v['subscriptionName'] == u['subscriptionName']
        assert sheet.cell(i, 2).value == v['resourceGroup'] == u['resourceGroup']
        assert sheet.cell(i, 3).value == u['name'] and sheet.cell(i, 8).value == u['status']
        expected = r['targetSKU'] if r['status'] == '建议试点调整' else r['currentSKU']
        assert u['effectiveSKU'] == expected and sheet.cell(i, 7).value == (expected if u['changing'] else '不更换')
        assert sheet.cell(i, 26).value == u['note']
        if r['kind'] == 'vmss':
            assert r['metricResourceId'].split('/')[-1] in u['note'] and '所有运行VM' in u['note']
        for col, key in zip(range(21, 26), ['payg', 'ri1', 'ri3', 'sp1', 'sp3']):
            cost, q = u['costs'][key], u['quotes'].get(key)
            display = sheet.cell(i, col).value
            # the shown text of a formula cell is only known after Excel recalculation
            pending = unrecalculated_formula(sheet.cell(i, col), recalculated)
            if u.get('deallocated', False):
                assert pending or display == ('0.00（省 0.00）' if key == 'payg' else '不建议购买')
                continue
            if cost is None:
                assert display in ('无报价', '暂不判断')
                continue
            validate_quote(q, key)
            if q is not None:
                same(cost, q['retailPrice'] / (12 if key == 'ri1' else 36) if key.startswith('ri') else q['retailPrice'] * 730)
            if not pending and u['currentMonthlyPAYG'] is not None:
                saving = u['currentMonthlyPAYG'] - cost
                assert display == f'{shown(cost)}（{"省" if saving >= 0 else "增加"} {shown(abs(saving))}）', (u['name'], key, display)
    assert calc.sheet_state == 'hidden' and sheet.max_column == 27
    w.close()
    md = (reports / f'{STEM}.md').read_text(encoding='utf-8')
    assert '官网公开列表价' in md and '|订阅|资源组|' in md
    tables, inside, ts, columns = [], False, [], 0
    for t in MarkdownIt('commonmark').enable('table').parse(md):
        if t.type == 'table_open':
            inside, ts = True, []
        if inside and t.type == 'tr_open':
            columns = 0
        if inside and t.type in ('td_open', 'th_open'):
            columns += 1
        if inside and t.type == 'tr_close':
            ts.append(columns)
        if t.type == 'table_close':
            assert len(set(ts)) == 1, 'Markdown table with inconsistent columns'
            tables.append(len(ts))
            inside = False
    assert tables == [7, len(rows) + 1], tables
    prices = read_json(f / 'price-summary.json')
    baseline = math.fsum(u['currentMonthlyPAYG'] for u in rows if u['running'])
    same(prices['currentMonthly'], baseline)
    for key in ['payg', 'ri1', 'ri3', 'sp1', 'sp3']:
        value = math.fsum(u['costs'][key] if u['costs'][key] is not None else u['costs']['payg'] for u in rows if u['running'])
        assert f'{value:,.2f}' in md
    result = {'status': 'success', 'subscriptionCount': len(subs), 'subscriptionWithVMCount': len({v['subscriptionId'] for v in vs}), 'instanceCount': len(vs),
              'runningCount': sum(r['state'] == 'PowerState/running' for r in rs), 'deallocatedCount': sum(r['state'] == 'PowerState/deallocated' for r in rs),
              'recommendedCount': sum(r['status'] == '建议试点调整' for r in rs), 'recommendedPoolCount': sum(p['status'] == '建议整池试点' for p in ps),
              'monthlySavingsUSD': total, 'formulaCount': verification['formulaCount'], 'recalculatedInExcel': verification['recalculatedInExcel'],
              'statusCounts': dict(Counter(u['status'] for u in rows))}
    write_json(f / 'audit.json', result)
    log(f'02 核对通过：{result["instanceCount"]} 台机器、{result["recommendedCount"]} 台建议调整、每月节省 USD {total:,.2f}，公式和表格与数据一致')
    return result
