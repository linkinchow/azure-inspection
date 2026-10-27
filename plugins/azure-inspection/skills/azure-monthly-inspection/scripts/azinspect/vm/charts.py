"""Charts for the VM report: PNG for Markdown and native Excel charts."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from matplotlib.ticker import FuncFormatter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.marker import DataPoint

from ..core import charts as cs
from .wording import STATUS

BUCKETS = [(0, 10, '<10%'), (10, 20, '10–20%'), (20, 40, '20–40%'), (40, 60, '40–60%'), (60, 80, '60–80%'), (80, 101, '≥80%')]
FILES = {'plans': '02_1_VM各方案每月费用.png', 'distribution': '02_2_VM平均利用率分布.png',
         'status': '02_3_VM建议结果.png', 'subscriptions': '02_4_各订阅每月预计节省.png'}
FILE_PATHS = {k: f'charts/{v}' for k, v in FILES.items()}


def chart_data(vms, recs, portfolio):
    running = [v for v in vms if v['state'] == 'PowerState/running' and v['cpuAverage'] is not None and v['memoryAverage'] is not None]
    dist = [(label, sum(lo <= v['cpuAverage'] < hi for v in running), sum(lo <= v['memoryAverage'] < hi for v in running)) for lo, hi, label in BUCKETS]
    status = Counter(STATUS.get(r['status'], r['status']) for r in recs)
    savings = Counter()
    for r in recs:
        savings[r['subscriptionName']] += r.get('monthlySavingsUSD') or 0
    plans = {p['key']: p for p in portfolio['plans']}
    return {'runningWithData': len(running), 'distribution': dist,
            'cpuBelow20': sum(v['cpuAverage'] < 20 for v in running), 'memoryBelow20': sum(v['memoryAverage'] < 20 for v in running),
            'status': sorted(status.items(), key=lambda kv: -kv[1]),
            'subscriptions': [(k, v) for k, v in savings.most_common() if v > 0],
            'plans': [('当前规格，按需', portfolio['currentMonthly'])] + [(name, plans[key]['monthlyCost']) for key, name in
                      [('payg', '降配后，按需'), ('sp1', '降配后，1 年 SP'), ('ri1', '降配后，1 年 RI'), ('sp3', '降配后，3 年 SP'), ('ri3', '降配后，3 年 RI')]]}


def make_pngs(data, chart_dir: Path):
    fig, ax = cs.figure(10, 3.8)
    cs.style(ax, 'y')
    x, w = range(len(data['distribution'])), 0.36
    ax.bar([i - w / 2 - 0.02 for i in x], [c for _, c, _ in data['distribution']], width=w, color=cs.BLUE, label='CPU')
    ax.bar([i + w / 2 + 0.02 for i in x], [m for _, _, m in data['distribution']], width=w, color=cs.ORANGE, label='内存')
    for i, (_, c, m) in enumerate(data['distribution']):
        ax.text(i - w / 2 - 0.02, c + 2, str(c), ha='center', fontsize=8.5, color=cs.INK)
        ax.text(i + w / 2 + 0.02, m + 2, str(m), ha='center', fontsize=8.5, color=cs.INK)
    ax.set_xticks(list(x), [b for b, _, _ in data['distribution']], fontsize=9, color=cs.INK)
    ax.set_xlabel('7 天平均利用率区间', fontsize=9, color=cs.INK2)
    ax.set_ylabel('VM 台数', fontsize=9, color=cs.INK2)
    ax.set_title(f'运行中 VM 的 CPU 与内存平均利用率分布（{data["runningWithData"]} 台有数据）', loc='left', fontsize=11, color=cs.INK, pad=10)
    ax.legend(frameon=False, fontsize=9, labelcolor=cs.INK2)
    cs.save(fig, chart_dir / FILES['distribution'])

    rows = list(reversed(data['status']))
    fig, ax = cs.figure(10, 0.4 * len(rows) + 1.2)
    cs.style(ax, 'x')
    ax.barh(range(len(rows)), [v for _, v in rows], height=0.56, color=[cs.BLUE if k.startswith('建议调整') else cs.GRAY for k, _ in rows])
    for i, (_, v) in enumerate(rows):
        ax.text(v + 1.5, i, str(v), va='center', fontsize=9, color=cs.INK)
    ax.set_yticks(range(len(rows)), [k for k, _ in rows], fontsize=9, color=cs.INK)
    ax.set_xlim(0, max(v for _, v in rows) * 1.12)
    ax.set_title(f'{sum(v for _, v in rows)} 台 VM 的建议结果', loc='left', fontsize=11, color=cs.INK, pad=10)
    cs.save(fig, chart_dir / FILES['status'])

    rows = list(reversed(data['subscriptions']))
    if rows:
        fig, ax = cs.figure(10, 0.4 * len(rows) + 1.2)
        cs.style(ax, 'x')
        ax.barh(range(len(rows)), [v for _, v in rows], height=0.56, color=cs.BLUE)
        total = sum(v for _, v in rows)
        for i, (_, v) in enumerate(rows):
            ax.text(v + total * 0.004, i, f'{v:,.0f}（{v / total:.0%}）', va='center', fontsize=9, color=cs.INK)
        ax.set_yticks(range(len(rows)), [k for k, _ in rows], fontsize=9, color=cs.INK)
        ax.set_xlim(0, max(v for _, v in rows) * 1.22)
        ax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: f'{x / 1000:,.0f}K' if x >= 1000 else f'{x:,.0f}'))
        ax.set_title(f'各订阅每月预计节省（合计 USD {total:,.2f}，官网列表价）', loc='left', fontsize=11, color=cs.INK, pad=10)
        cs.save(fig, chart_dir / FILES['subscriptions'])

    rows, base = list(reversed(data['plans'])), data['plans'][0][1]
    fig, ax = cs.figure(10, 3.6)
    cs.style(ax, 'x')
    for i, (name, v) in enumerate(rows):
        ax.barh(i, v, height=0.55, color=cs.GRAY if name.startswith('当前') else cs.BLUE)
        ax.text(v + base * 0.01, i, f'{v:,.0f}' if name.startswith('当前') else f'{v:,.0f}（每月省 {base - v:,.0f}）', va='center', fontsize=9, color=cs.INK)
    ax.set_yticks(range(len(rows)), [n for n, _ in rows], fontsize=10, color=cs.INK)
    ax.set_xlim(0, base * 1.35)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: f'{x / 1000:,.0f}K'))
    ax.set_title('VM 每月费用：当前与各方案对比（USD，官网列表价）', loc='left', fontsize=11, color=cs.INK, pad=10)
    cs.save(fig, chart_dir / FILES['plans'])
    return {k: v for k, v in FILE_PATHS.items() if (chart_dir / FILES[k]).exists()}


def add_excel_sheet(sheet, data):
    ws = sheet('图表', 'VM 利用率与节省图表', '图表数据来自本工作簿；价格为Azure官网公开列表价，不是合同价或实际账单价。', ['项目', '值1', '值2'])
    blocks = {}

    def block(key, header, rows):
        if not rows:
            return
        start = ws.max_row + 2
        for i, h in enumerate(header):
            ws.cell(start, i + 1, h)
        for r, row in enumerate(rows, start + 1):
            for c, v in enumerate(row, 1):
                ws.cell(r, c, v)
        blocks[key] = (start, start + len(rows))
    block('distribution', ['利用率区间', 'CPU台数', '内存台数'], data['distribution'])
    block('status', ['建议结果', 'VM台数'], data['status'])
    block('subscriptions', ['订阅', '每月预计节省USD'], data['subscriptions'])
    block('plans', ['方案', '每月费用USD'], data['plans'])

    def chart(key, title, anchor, cols, colors, horizontal, number_format='#,##0', highlight=None):
        if key not in blocks:
            return
        first, last = blocks[key]
        ch = BarChart()
        ch.type = 'bar' if horizontal else 'col'
        ch.height, ch.width, ch.gapWidth = 7.5, 18, 60
        ch.y_axis.numFmt = number_format
        for col, color in zip(cols, colors):
            ch.add_data(Reference(ws, min_col=col, min_row=first, max_row=last), titles_from_data=True)
            ch.series[-1].graphicalProperties.solidFill = color
            ch.series[-1].graphicalProperties.line.solidFill = color
        ch.set_categories(Reference(ws, min_col=1, min_row=first + 1, max_row=last))
        if highlight:
            for idx in range(last - first):
                pt = DataPoint(idx=idx)
                fill = highlight(ws.cell(first + 1 + idx, 1).value)
                pt.graphicalProperties.solidFill = fill
                pt.graphicalProperties.line.solidFill = fill
                ch.series[0].dPt.append(pt)
        if horizontal:
            ch.x_axis.scaling.orientation = 'maxMin'
            ch.y_axis.crosses = 'max'
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        if len(cols) > 1:
            ch.legend.position, ch.legend.overlay = 'b', False
        else:
            ch.legend = None
        ws[f'{anchor[0]}{int(anchor[1:]) - 1}'] = title
        ws.add_chart(ch, anchor)

    chart('distribution', 'CPU 与内存平均利用率分布（VM 台数）', 'F5', [2, 3], ['2A78D6', 'EB6834'], False, '0')
    chart('status', '各类建议结果的 VM 台数', 'F22', [2], ['A09F98'], True, '0', lambda v: '2A78D6' if str(v).startswith('建议调整') else 'A09F98')
    chart('subscriptions', '各订阅每月预计节省（USD，列表价）', 'F39', [2], ['2A78D6'], True)
    chart('plans', 'VM 每月费用：当前与各方案（USD，列表价）', 'F56', [2], ['2A78D6'], True, '#,##0', lambda v: 'A09F98' if str(v).startswith('当前') else '2A78D6')
    return ws
