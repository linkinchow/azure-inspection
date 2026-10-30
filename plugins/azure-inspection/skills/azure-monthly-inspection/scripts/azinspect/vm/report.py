"""02 虚拟机利用率与规格优化: Excel (with formulas verified) and Markdown."""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties

from ..core.charts import clear as clear_charts
from ..core.excel import MONEY, verify
from ..core.fmt_md import format_and_check
from ..core.io import write_json
from . import charts as vm_charts
from . import scalesets
from .prices import KEYS
from .rules import IO
from .table import combined_rows, portfolio_totals, price_and_saving

STEM = '02_虚拟机利用率与规格优化'
DETAIL_SHEET = 'VM建议指标价格总表'
PRICE_NOTE = '所有价格均为Azure官网公开列表价（USD），包括按需、RI和SP；不是客户合同价或实际账单价。节省也按列表价计算，未考虑客户已有折扣。'


def fmt(value):
    return '无数据' if value is None else f'{value:,.2f}'


def esc(value):
    return str(value).replace('|', '/').replace('\n', ' ')


def _sum_rows(rows):
    """SUM of column D on '价格计算依据' over contiguous row ranges, at most 200 ranges per SUM (Excel allows 255 arguments).
    Adding one cell per row exceeds Excel's 8,192-character formula limit for subscriptions with thousands of machines."""
    ranges = []
    for r in rows:
        if ranges and r == ranges[-1][1] + 1:
            ranges[-1][1] = r
        else:
            ranges.append([r, r])
    refs = [f"'价格计算依据'!D{a}" if a == b else f"'价格计算依据'!D{a}:D{b}" for a, b in ranges]
    return '=' + '+'.join('SUM(' + ','.join(refs[k:k + 200]) + ')' for k in range(0, len(refs), 200))


def _unverified_reason(sub):
    errors = sub.get('inventoryErrors') or []
    if sub.get('state') not in (None, 'Enabled') or any(e.get('phase') == 'subscriptionState' for e in errors):
        return '订阅未启用，无法读取'
    return '读取失败或在 Resource Graph 中不可见，无法确认资源'


def _unverified_text(unverified):
    """'另有 17 个订阅无法确认资源：订阅未启用，无法读取（a、b、c 等 17 个）。' grouped by reason."""
    if not unverified:
        return ''
    groups = {}
    for s in unverified:
        groups.setdefault(_unverified_reason(s), []).append(str(s['name']))
    parts = []
    for reason, names in groups.items():
        names = sorted(names)
        listed = '、'.join(names[:3]) + (f' 等 {len(names)} 个' if len(names) > 3 else '')
        parts.append(f'{reason}（{listed}）')
    return f'另有{len(unverified)}个订阅无法确认资源：' + '；'.join(parts) + '。'


def build(ws, cfg, narrative=None, log=print):
    narrative = narrative or {}
    fdir, out = ws.month_dir('findings') / 'vm', ws.month_dir('reports')
    read = lambda name: __import__('json').loads((fdir / name).read_text(encoding='utf-8'))
    vms, recs, pools, manifest = read('vms-analyzed.json'), read('recommendations.json'), read('pools.json'), read('collection-manifest.json')
    subscriptions = read('subscriptions.json')
    from .prices import price_catalog
    catalog = price_catalog(ws.raw('vm/prices'), read('summary.json')['priceAsOfDate'])
    from ..cost.commitments import load as load_commitments, vm_coverage
    coverage, commitments = vm_coverage(ws), load_commitments(ws)
    month = (commitments or {}).get('month')
    unified = combined_rows(vms, recs, pools, catalog, narrative.get('crossCheckNotes'), coverage,
                            f'{int(month[5:])}月' if month else '巡检月份')
    covered = [u for u in unified if u['changing'] and (u.get('committedShare') or 0) >= 0.5]
    cc_info = {
        'read': coverage is not None, 'attempted': commitments is not None, 'month': month, 'coveredCount': len(covered),
        'coveredSavingsUSD': sum(u['monthlyRightsizingSavings'] or 0 for u in covered),
        'machines': [{'name': u['name'], 'subscription': u['subscriptionName'], 'share': u['committedShare']} for u in covered]}
    write_json(fdir / 'commitment-coverage.json', cc_info)
    # subscription ID after the name keeps each subscription's rows contiguous even when two subscriptions share a name
    unified.sort(key=lambda r: (str(r['subscriptionName']).lower(), r['subscriptionId'].lower(), str(r['resourceGroup']).lower(), r['name'].lower(),
                                r['vmId'].lower()))
    portfolio = portfolio_totals(unified)
    write_json(fdir / 'rows.json', unified)
    write_json(fdir / 'price-summary.json', portfolio)
    wb, expected = Workbook(), {}
    wb.remove(wb.active)

    def sheet(name, title, note, columns):
        s = wb.create_sheet(name)
        s.append([title])
        s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        s.append([note])
        s.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
        s.append(columns)
        s.freeze_panes = 'A4'
        return s

    def formula(s, row, col, expression, value):
        s.cell(row, col, expression)
        expected[f'{s.title}!{L(col)}{row}'] = value

    count_states = Counter(r['state'] for r in recs)
    accepted = [r for r in recs if r['status'] == '建议试点调整']
    total = sum(r['monthlySavingsUSD'] for r in recs)
    scopes = {s['id']: [u for u in unified if u['subscriptionId'] == s['id']] for s in subscriptions}
    unverified = [s for s in subscriptions if s.get('inventoryStatus') != 'ok']
    assert sum(map(len, scopes.values())) == len(unified)
    overview = sheet('总览', f'{cfg.name}：跨订阅VM规格建议与预计节省', PRICE_NOTE, ['项目', '结果', '说明'])
    coverage_note = f'{cfg.name}纳入巡检的全部订阅' + (f'；另有{len(unverified)}个订阅无法确认资源，见“订阅范围”页' if unverified else '')
    for row in [
        ['检查订阅数', len(subscriptions) - len(unverified), coverage_note],
        ['有VM的订阅数', sum(bool(x) for x in scopes.values()), '没有VM的订阅也在订阅范围页列出'],
        ['观察窗口', manifest['start'] + ' 至 ' + manifest['end'], '全部使用同一完整7天，日平均按样本数加权'],
        ['VM/VMSS实例数', len(unified), '总表最前两列为订阅、资源组'],
        ['运行实例数', count_states['PowerState/running'], '已释放与状态不明分别处理'],
        ['已释放实例数', count_states['PowerState/deallocated'], '已停止计算收费，不购买持续计算承诺'],
        ['状态不明或其他状态', len(recs) - count_states['PowerState/running'] - count_states['PowerState/deallocated'], '不能当成已释放或空闲'],
        ['建议调整机器数', len(accepted), '平均用量、磁盘、规格价格及同池规则满足；先小范围试用'],
        ['预计每月节省USD', None, '仅计算规格调整；RI/SP价格另列'],
        ['预计每年节省USD', None, '每月预计节省×12'],
        ['数据或权限缺失', '见总表备注与订阅范围', '缺数据不补0；不能判断时不推荐降配'],
    ]:
        overview.append(row)
    calculation = sheet('价格计算依据', '价格与每月节省的计算依据', PRICE_NOTE,
                        ['VM资源ID', '当前规格按需月USD', '调整后按需月USD', '按需每月节省USD', '调整后1年RI月USD', '1年RI每月节省USD', '调整后3年RI月USD',
                         '3年RI每月节省USD', '调整后1年SP月USD', '1年SP每月节省USD', '调整后3年SP月USD', '3年SP每月节省USD'])
    for j, u in enumerate(unified, 4):
        calculation.append([u['vmId'], u['currentMonthlyPAYG']] + [None] * 10)
        for key, col in zip(KEYS, [3, 5, 7, 9, 11]):
            cost, quote = u['costs'][key], u['quotes'].get(key)
            if u.get('deallocated', False):
                calculation.cell(j, col, 0 if key == 'payg' else '不建议购买')
            elif cost is None:
                calculation.cell(j, col, '无报价/暂不判断')
            elif quote is None:
                calculation.cell(j, col, cost)
            else:
                divisor = 12 if key == 'ri1' else 36 if key == 'ri3' else None
                expression = f'={quote["retailPrice"]!r}/{divisor}' if divisor else f'={quote["retailPrice"]!r}*730'
                formula(calculation, j, col, expression, cost)
            if cost is not None and u['currentMonthlyPAYG'] is not None:
                formula(calculation, j, col + 1, f'=B{j}-{L(col)}{j}', u['currentMonthlyPAYG'] - cost)
    for col, key in enumerate(KEYS, 13):
        calculation.cell(3, col, '汇总使用的' + key + '月USD')
    for j, u in enumerate(unified, 4):
        for col, key, source in zip(range(13, 18), KEYS, ['C', 'E', 'G', 'I', 'K']):
            if not u['running']:
                calculation.cell(j, col, 0)
                continue
            value = u['costs'][key] if u['costs'][key] is not None else u['costs']['payg']
            formula(calculation, j, col, f'=IF(ISNUMBER({source}{j}),{source}{j},C{j})', value)
    comparison = sheet('费用汇总', '当前总价与调整后五种方案', PRICE_NOTE + 'RI缺报价时按调整后按需价计入；所有节省都与当前规格按需列表价比较。',
                       ['方案', '每月总价USD', '每月节省USD', '每年总价USD', '每年节省USD', '说明'])
    comparison.append(['当前规格：按需', None, None, None, None, '当前运行机器按每月730小时计算；已释放机器计算费为0'])
    for plan in portfolio['plans']:
        comparison.append([plan['label'], None, None, None, None, plan['note']])
    calculation.sheet_state = 'hidden'
    headers = ['订阅', '资源组', 'VM/机器', 'Region', '状态', '当前规格', '建议更换规格', '建议怎么处理',
               '更换前CPU平均%', '更换后CPU平均%（预计）', '更换前内存平均%', '更换后内存平均%（预计）',
               '更换前Uncached IOPS平均%', '更换后Uncached IOPS平均%（预计）', '更换前Uncached吞吐平均%', '更换后Uncached吞吐平均%（预计）',
               '更换前Cached IOPS平均%', '更换后Cached IOPS平均%（预计）', '更换前Cached吞吐平均%', '更换后Cached吞吐平均%（预计）',
               '调整后规格：按需月价USD（每月节省）', '调整后规格：1年RI月价USD（每月节省）', '调整后规格：3年RI月价USD（每月节省）',
               '调整后规格：1年SP月价USD（每月节省）', '调整后规格：3年SP月价USD（每月节省）', '备注（含VMSS/节点池）', 'VM资源ID']
    table = wb.create_sheet(DETAIL_SHEET)
    table.append([f'{cfg.name}：每台机器的规格、VM指标和价格'])
    table.merge_cells(start_row=1, start_column=1, end_row=1, end_column=27)
    table.append([PRICE_NOTE + ' 价格格=月价（比当前规格按需月价每月省多少）；不更换时按当前规格报价。'])
    table.merge_cells(start_row=2, start_column=1, end_row=2, end_column=27)
    for start, end, title in [(1, 8, '订阅、资源组和规格建议'), (9, 20, 'VM指标：更换前 → 更换后（预计）'),
                              (21, 25, '调整后规格公开列表月价USD（相对当前规格按需价的每月节省）'), (26, 27, '原因和一起调整的VMSS名称')]:
        table.cell(3, start, title)
        table.merge_cells(start_row=3, start_column=start, end_row=3, end_column=end)
    table.append(headers)
    table.freeze_panes = 'I5'
    table.sheet_properties.outlinePr.summaryRight = False
    for start, end in [(9, 20), (21, 25)]:
        table.column_dimensions.group(L(start), L(end), outline_level=1, hidden=False)
    table.column_dimensions['AA'].hidden = True
    for i, u in enumerate(unified, 5):
        table.append([u['subscriptionName'], u['resourceGroup'], u['name'], u['region'], u['state'], u['currentSKU'],
                      u['effectiveSKU'] if u['changing'] else '不更换', u['status'], u['cpuCurrent'], u['cpuAfter'], u['memoryCurrent'], u['memoryAfter']]
                     + [val for metric, _key, _label in IO for val in (u['currentIO'].get(metric), u['afterIO'].get(metric))]
                     + [None] * 5 + [u['note'], u['vmId']])
        j = i - 1
        for col, key, numeric_col in zip(range(21, 26), KEYS, [3, 5, 7, 9, 11]):
            text = price_and_saving(u, key)
            if u['costs'][key] is None or u['currentMonthlyPAYG'] is None:
                table.cell(i, col, text)
            else:
                costref, savingref = f"'价格计算依据'!{L(numeric_col)}{j}", f"'价格计算依据'!{L(numeric_col + 1)}{j}"
                expression = f'=TEXT({costref},"#,##0.00")&"（"&IF({savingref}>=0,"省 ","增加 ")&TEXT(ABS({savingref}),"#,##0.00")&"）"'
                formula(table, i, col, expression, {'text': text})
    formula(overview, 12, 2, f"=SUM('价格计算依据'!D4:D{calculation.max_row})", total)
    formula(overview, 13, 2, '=B12*12', total * 12)
    formula(comparison, 4, 2, f"=SUM('价格计算依据'!B4:B{calculation.max_row})", portfolio['currentMonthly'])
    formula(comparison, 4, 3, '=0', 0)
    formula(comparison, 4, 4, '=B4*12', portfolio['currentAnnual'])
    formula(comparison, 4, 5, '=0', 0)
    for i, plan in enumerate(portfolio['plans'], 5):
        numeric_col = {'payg': 'M', 'ri1': 'N', 'ri3': 'O', 'sp1': 'P', 'sp3': 'Q'}[plan['key']]
        formula(comparison, i, 2, f"=SUM('价格计算依据'!{numeric_col}4:{numeric_col}{calculation.max_row})", plan['monthlyCost'])
        formula(comparison, i, 3, f'=$B$4-B{i}', plan['monthlySaving'])
        formula(comparison, i, 4, f'=B{i}*12', plan['annualCost'])
        formula(comparison, i, 5, f'=C{i}*12', plan['annualSaving'])
    scope = sheet('订阅范围', '每个订阅的覆盖情况和预计节省', '无VM不等于订阅无费用；缺指标/权限的机器仍保留，缺失不按0处理。',
                  ['订阅', 'VM实例数', '运行数', '已释放数', '其他/状态不明', '建议调整数', '预计每月节省USD', '说明', '订阅ID'])
    for i, s in enumerate(sorted(subscriptions, key=lambda s: str(s['name']).lower()), 4):
        us = scopes[s['id']]
        ids = {u['vmId'] for u in us}
        ss = [r for r in recs if r['vmId'] in ids]
        sc = Counter(r['state'] for r in ss)
        insufficient = sum('数据' in u['status'] or '权限' in u['note'] for u in us)
        note = (_unverified_reason(s) if s.get('inventoryStatus') != 'ok' else '已确认没有VM/VMSS；其他服务费用另行分析' if not us
                else f'{insufficient}台数据/权限待确认' if insufficient else '范围内机器已列入总表')
        scope.append([s['name'], len(us), sc['PowerState/running'], sc['PowerState/deallocated'], len(us) - sc['PowerState/running'] - sc['PowerState/deallocated'],
                      sum(u['changing'] for u in us), None, note, s['id']])
        expression = _sum_rows([j for j, u in enumerate(unified, 4) if u['subscriptionId'] == s['id']]) if us else '=0'
        formula(scope, i, 7, expression, sum(r['monthlySavingsUSD'] for r in ss))
    daily = sheet('每日指标原始值', '每日Average和有效样本数', '可用内存指标是可用率；总表内存使用率=100%−可用率。缺失保留空值。',
                  ['订阅', '资源组', 'VM/机器', '指标', '日桶开始UTC', 'Average', 'Count', 'VM资源ID'])
    for v in vms:
        for metric, m in v['metrics'].items():
            for d in m.get('daily', []):
                daily.append([v['subscriptionName'], v['resourceGroup'], v['metricName'], metric, d['timeStamp'], d.get('average'), d.get('count'), v['id']])
    method = sheet('规则与来源', '规则、列表价和调整前注意事项', PRICE_NOTE, ['项目', '说明'])
    for row in [
        ['CPU/内存', '准备减少的那项周平均须<20%；另一项容量保留。每项最多减半，同代保留原功能。'],
        ['磁盘', '目标Cached/Uncached IOPS和吞吐平均消耗率各≤50%，四项不相加，不查单盘IO或P99。'],
        ['数据完整性', '须7个有效日、相对采样数量≥95%。缺内存、缺磁盘路径或读权限时不推荐。'],
        ['SKU限制', '按机器所在订阅、区域和可用区检查目标规格是否可部署。'],
        ['VMSS/AKS', '同池运行节点须有共同且通过检查的目标，备注写VMSS名称及所有VM一起调整。'],
        ['指标箭头', '更换前是实测周平均；更换后是预计值。写不更换时两边相同。'],
        ['按需与SP月费用', '官方每小时列表价×730，按全年稳定用量估算。'],
        ['RI月费用', '1年订单公开列表价÷12；3年订单价÷36。未找到对应报价时显示无报价。'],
        ['停机与状态不明', '只有deallocated明确停计算收费。其他状态不当成停机，缺数据不购买承诺。'],
        ['变更前', '确认重要任务、故障接管、应用资源预留、库存和配额，先小范围调整；异常时恢复原配置。'],
        ['Azure VM官网定价入口', 'https://azure.microsoft.com/en-us/pricing/details/virtual-machines/linux/'],
        ['Batch API', 'https://learn.microsoft.com/en-us/rest/api/monitor/metrics-batch/batch?view=rest-monitor-2023-10-01'],
        ['磁盘指标', 'https://learn.microsoft.com/en-us/azure/virtual-machines/disks-metrics'],
    ]:
        method.append(row)
    data = vm_charts.chart_data(vms, recs, portfolio)
    data['commitmentCoverage'] = cc_info
    analysis = read('summary.json')
    data['scaleSetNote'] = scalesets.note(analysis.get('flexibleDuplicatesRemoved', 0), analysis.get('flexibleInstancesMissing', 0))
    clear_charts(out / 'charts', '02')
    pngs = vm_charts.make_pngs(data, out / 'charts')
    chart_sheet = vm_charts.add_excel_sheet(sheet, data)
    for s in wb:
        s.sheet_view.showGridLines = False
        s.auto_filter.ref = f'A3:{L(s.max_column)}{s.max_row}'
        s.print_title_rows = '1:3'
        s.page_setup.orientation = 'landscape'
        s.page_setup.fitToWidth = 1
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.row_dimensions[1].height, s.row_dimensions[2].height, s.row_dimensions[3].height = 28, 48, 48
        for row in s:
            for c in row:
                c.font = Font(name='Arial', size=10)
                c.alignment = Alignment(vertical='top', wrap_text=True)
                if c.data_type == 'f':
                    c.font = Font(name='Arial', size=10, color='008000' if '!' in c.value else '000000')
                if isinstance(c.value, (int, float)) or c.data_type == 'f':
                    c.number_format = MONEY
                if c.row == 1:
                    c.fill = PatternFill('solid', fgColor='17365D')
                    c.font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
                elif c.row == 3:
                    c.fill = PatternFill('solid', fgColor='DCE6F1')
                    c.font = Font(name='Arial', size=10, bold=True)
                elif c.row > 3 and c.row % 2 == 0:
                    c.fill = PatternFill('solid', fgColor='F5F8FC')
                if isinstance(c.value, str) and c.value.startswith('https://'):
                    c.hyperlink = c.value
        for col in range(1, s.max_column + 1):
            s.column_dimensions[L(col)].width = 24
        s.column_dimensions['A'].width = 32
    table.auto_filter.ref = f'A4:AA{table.max_row}'
    table.print_title_rows = '1:4'
    table.row_dimensions[3].height, table.row_dimensions[4].height = 30, 70
    for start, end, color in [(1, 8, 'DCE6F1'), (9, 20, 'E2F0D9'), (21, 25, 'FFF2CC'), (26, 27, 'E4DFEC')]:
        for rr in [3, 4]:
            for col in range(start, end + 1):
                c = table.cell(rr, col)
                c.fill = PatternFill('solid', fgColor=color)
                c.font = Font(name='Arial', size=10, bold=True)
    for col, width in [('B', 42), ('C', 44), ('H', 44), ('Z', 110), ('AA', 100)]:
        table.column_dimensions[col].width = width
    for i in range(5, table.max_row + 1):
        table.row_dimensions[i].height = 85
    overview.column_dimensions['B'].width, overview.column_dimensions['C'].width = 66, 95
    for i in range(4, overview.max_row + 1):
        overview.row_dimensions[i].height = 42
    scope.column_dimensions['H'].width, scope.column_dimensions['I'].width = 75, 44
    daily.column_dimensions['D'].width, daily.column_dimensions['H'].width = 56, 100
    method.column_dimensions['B'].width = 120
    wb._sheets.remove(comparison)
    wb._sheets.insert(1, comparison)
    wb._sheets.remove(chart_sheet)
    wb._sheets.insert(2, chart_sheet)
    chart_sheet.column_dimensions['A'].width = 30
    comparison.column_dimensions['A'].width, comparison.column_dimensions['F'].width = 38, 88
    for i in range(4, comparison.max_row + 1):
        comparison.row_dimensions[i].height = 48
    out.mkdir(parents=True, exist_ok=True)
    path = out / f'{STEM}.xlsx'
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
    write_json(fdir / 'expected-formulas.json', expected)
    md_path = out / f'{STEM}.md'
    md_path.write_text(markdown(cfg, unified, recs, subscriptions, manifest, pools, total, portfolio, data, pngs, narrative), encoding='utf-8')
    result = verify(path, expected, fdir / 'verification.json')
    log(f'02 报告：{len(unified)} 台机器，建议调整 {len(accepted)} 台，预计每月节省 USD {total:,.2f}（官网列表价）；'
        f'Excel 公式 {result["formulaCount"]} 个' + ('，已用 Excel 复算核对' if result['recalculatedInExcel'] else f'，本机未能用 Excel 复算（{result["recalculationError"]}）'))
    return {'markdown': md_path, 'excel': path, 'monthlySavingsUSD': total, 'recommended': len(accepted), 'instances': len(unified)}


def _result(u):
    """Outcome for the summary; '暂不调整：见备注' is split by the first reason in the note."""
    if u['status'] == '暂不调整：见备注':
        reason = (u['note'] or '').split('；')[0] or '见备注'
        if '平均消耗率' in reason:   # one row per measured percentage would split the table: group the disk limits
            reason = '更换后磁盘 IOPS 或吞吐的平均消耗率超过 50%'
        return '暂不调整：' + reason
    return u['status']


def _change(a, b):
    return fmt(a) + ' → ' + fmt(b)


def _mean(values):
    values = [v for v in values if v is not None]
    return sum(values) / len(values) if values else None


def detail_table(rows):
    """The per-machine table (same content as the Excel sheet), attached only up to cfg.vm_detail_max_rows machines."""
    def disk(u, start):
        return '；'.join(label + ' ' + _change(u['currentIO'].get(IO[i][0]), u['afterIO'].get(IO[i][0])) for i, label in [(start, 'IOPS'), (start + 1, '吞吐')])
    lines = ['- **排序**：依次按订阅名、资源组名、VM名称的字母顺序排列，忽略大小写。',
             '- **资源定位**：按订阅、资源组和机器名称查找资源；完整资源ID见Excel。',
             '- **建议更换规格**：建议调整的写新SKU，其余写“不更换”。不更换按当前规格计算指标和价格。',
             '- **指标箭头**：更换前实测周平均 → 更换后预计周平均；不更换时两边相同。Cached/Uncached各自显示IOPS和吞吐消耗百分比。',
             '- **五列价格**：均按调整后规格报价；写“不更换”的仍按当前规格。每格为“每台月价（每月节省）”，单位USD。括号里的节省统一与当前规格按需列表价比较，已包含规格调整和所选RI/SP优惠。各列是互斥方案，不能相加。',
             '- **月价计算**：按需/SP每小时列表价×730，1年RI订单价÷12，3年RI订单价÷36。未找到报价显示“无报价”；已停机机器不建议购买RI/SP。',
             '- **备注**：说明原因、VMSS或节点池名称；同VMSS所有VM一起调整，不单独改某一个节点。', '',
             '|订阅|资源组|VM/机器|Region|当前规格|建议更换规格|建议|CPU平均%：前→后（预计）|内存平均%：前→后（预计）|Uncached IOPS/吞吐%：前→后（预计）|Cached IOPS/吞吐%：前→后（预计）|调整后规格：按需月价（每月节省）|调整后规格：1年RI月价（每月节省）|调整后规格：3年RI月价（每月节省）|调整后规格：1年SP月价（每月节省）|调整后规格：3年SP月价（每月节省）|备注（含VMSS）|',
             '|---|---|---|---|---|---|---|---|---|---|---|---:|---:|---:|---:|---:|---|']
    for u in rows:
        values = [u['subscriptionName'], u['resourceGroup'], u['name'], u['region'], u['currentSKU'], u['effectiveSKU'] if u['changing'] else '不更换', u['status'],
                  _change(u['cpuCurrent'], u['cpuAfter']), _change(u['memoryCurrent'], u['memoryAfter']), disk(u, 0), disk(u, 2)] + [price_and_saving(u, k) for k in KEYS] + [u['note']]
        lines.append('|' + '|'.join(esc(v) for v in values) + '|')
    return lines


def summary_section(rows, total, limit):
    """Section 2: four summary tables of fixed size, then the per-machine table when there are at most `limit` machines.
    A per-machine table for a large tenant runs to megabytes and duplicates the Excel sheet."""
    changing = [u for u in rows if u['changing']]
    saving = lambda us: sum(u['monthlyRightsizingSavings'] or 0 for u in us)  # noqa: E731
    lines = ['## 2. 建议汇总', '',
             f'下面四张表汇总 {len(rows)} 台机器的建议；每台机器的指标、五种价格和备注见 Excel“{DETAIL_SHEET}”，可按订阅、资源组和建议结果筛选。'
             '表中节省均为调整规格后、按需列表价每月节省的金额（USD）。', '',
             '### 2.1 按订阅汇总', '', '|订阅|机器数|运行中|建议调整|每月预计节省USD|', '|---|---:|---:|---:|---:|']
    by_sub = defaultdict(list)
    for u in rows:
        by_sub[(u['subscriptionName'], u['subscriptionId'])].append(u)
    for (name, _sid), us in sorted(by_sub.items(), key=lambda kv: (-saving(kv[1]), str(kv[0][0]).lower())):
        lines.append(f'|{esc(name)}|{len(us)}|{sum(u["running"] for u in us)}|{sum(u["changing"] for u in us)}|{fmt(saving(us))}|')
    lines += [f'|合计|{len(rows)}|{sum(u["running"] for u in rows)}|{len(changing)}|{fmt(total)}|', '',
              '### 2.2 按建议结果汇总', '', '|建议结果|台数|运行中|每月预计节省USD|', '|---|---:|---:|---:|']
    by_result = defaultdict(list)
    for u in rows:
        by_result[_result(u)].append(u)
    for result, us in sorted(by_result.items(), key=lambda kv: (not kv[1][0]['changing'], -len(kv[1]), kv[0])):
        lines.append(f'|{esc(result)}|{len(us)}|{sum(u["running"] for u in us)}|{fmt(saving(us))}|')
    if changing:
        by_sku = defaultdict(list)
        for u in changing:
            by_sku[(u['currentSKU'], u['effectiveSKU'])].append(u)
        top = sorted(by_sku.items(), key=lambda kv: (-saving(kv[1]), kv[0]))[:15]
        lines += ['', '### 2.3 按规格变化汇总', '', '|当前规格|建议规格|台数|每月预计节省USD|', '|---|---|---:|---:|']
        lines += [f'|{a}|{b}|{len(us)}|{fmt(saving(us))}|' for (a, b), us in top]
        if len(by_sku) > len(top):
            lines += ['', f'共 {len(by_sku)} 种规格变化，以上是节省最多的 {len(top)} 种。']
        items = defaultdict(list)
        for u in changing:
            items[(u['subscriptionName'], u['poolName']) if u['poolName'] else (u['subscriptionName'], u['resourceGroup'], u['name'])].append(u)
        best = sorted(items.values(), key=lambda us: (-saving(us), us[0]['name']))[:20]
        lines += ['', '### 2.4 节省最多的前 20 项', '', '节点池和 VMSS 的机器需要一起调整，合并为一行。', '',
                  '|订阅|资源组|VM/节点池|台数|当前规格 → 建议规格|CPU平均%：前→后（预计）|内存平均%：前→后（预计）|每月预计节省USD|备注|',
                  '|---|---|---|---:|---|---|---|---:|---|']
        for us in best:
            u = us[0]
            name = f'节点池/VMSS {u["poolName"]}' if u['poolName'] else u['name']
            skus = '、'.join(sorted({f'{x["currentSKU"]} → {x["effectiveSKU"]}' for x in us}))
            values = [u['subscriptionName'], '、'.join(sorted({x['resourceGroup'] for x in us})), name, len(us), skus,
                      _change(_mean(x['cpuCurrent'] for x in us), _mean(x['cpuAfter'] for x in us)),
                      _change(_mean(x['memoryCurrent'] for x in us), _mean(x['memoryAfter'] for x in us)), fmt(saving(us)), u['note']]
            lines.append('|' + '|'.join(esc(v) for v in values) + '|')
    if len(rows) <= limit:
        lines += ['', '### 2.5 逐台明细', ''] + detail_table(rows)
    else:
        lines += ['', f'共 {len(rows)} 台，逐台明细见 Excel“{DETAIL_SHEET}”，可按订阅、资源组、建议结果筛选。']
    return lines


def markdown(cfg, rows, recs, subs, manifest, pools, total, portfolio, data, pngs, narrative):
    accepted = [r for r in recs if r['status'] == '建议试点调整']
    state = Counter(r['state'] for r in recs)
    with_vm = {u['subscriptionId'] for u in rows}
    without = [s['name'] for s in subs if s['id'] not in with_vm and s.get('inventoryStatus') == 'ok']
    unverified = [s for s in subs if s.get('inventoryStatus') != 'ok']
    pool_nodes = sum(bool(r.get('pool')) for r in accepted)
    pool_count = sum(p['status'] == '建议整池试点' for p in pools)
    running, deallocated = state['PowerState/running'], state['PowerState/deallocated']
    fallback = [u for u in rows if u.get('metricsSource')]
    unverified_text = _unverified_text(unverified)
    lines = [f'# {cfg.name}：各订阅VM规格建议与预计节省', '', '**' + PRICE_NOTE + '**', '', '## 1. 范围和结果', '',
             f'- 已核对{cfg.name}纳入巡检的 **{len(subs) - len(unverified)}个订阅**：{len(with_vm)}个有VM/VMSS，{len(without)}个没有VM。{unverified_text}'
             f'发现的 **{len(rows)}台机器**中，{running}台运行、{deallocated}台已释放计算资源。',
             *([f'- {data["scaleSetNote"]}'] if data.get('scaleSetNote') else []),
             f'- 数据窗口为 **{manifest["start"]}至{manifest["end"]}（UTC）**，所有订阅使用同一完整7天。',
             f'- 建议调整 **{len(accepted)}台机器**，其中{len(accepted) - pool_nodes}台独立VM，{pool_nodes}台在{pool_count}个节点池中。节点池里的机器已包含在总数中。',
             f'- 仅调整规格、不购买RI/SP，**预计每月节省 USD {fmt(total)}，每年节省 USD {fmt(total * 12)}**。当前费用及五种方案汇总如下。',
             '- 没有数据或权限不足的机器仍列在逐台明细中，明确写“暂不判断”，不把缺数据当成空闲；数量见 2.2。', '',
             '**先小范围调整，确认业务正常后再推广。**平均用量低不代表每个忙碌时刻都够用；业务团队需确认重要任务、故障接管和AKS资源需求。报告提供分析建议，实际调整需经业务负责人批准。', '',
             '### 当前费用与调整后五种方案', '',
             f'**以下均按Azure官网公开列表价计算，单位USD，不是实际账单。**比较的是同一批{running}台运行机器的基础计算费用；{deallocated}台已释放机器计算费为0。每月按730小时估算。', '',
             '|方案|每月总价USD|每月节省USD|每年总价USD|每年节省USD|说明|', '|---|---:|---:|---:|---:|---|',
             f'|当前规格：按需|{fmt(portfolio["currentMonthly"])}|0.00|{fmt(portfolio["currentAnnual"])}|0.00|当前运行机器的按需列表价|']
    for plan in portfolio['plans']:
        lines.append('|' + '|'.join(esc(v) for v in [plan['label'], fmt(plan['monthlyCost']), fmt(plan['monthlySaving']), fmt(plan['annualCost']), fmt(plan['annualSaving']), plan['note']]) + '|')
    ri1 = next(p for p in portfolio['plans'] if p['key'] == 'ri1')
    lines += ['', (f'RI方案中，{ri1["fallbackCount"]}台机器没有对应RI报价，这部分仍按调整后规格的按需价计入，每月USD {fmt(ri1["fallbackMonthly"])}；逐台明细中仍显示“无报价”。'
                   if ri1['fallbackCount'] else '') + '所有方案的节省均与当前规格按需列表价比较。RI/SP方案已包含规格调整的效果，各方案分别选择，不相加。', '']
    for key, alt in [('plans', 'VM 各方案每月费用'), ('distribution', 'VM 平均利用率分布')]:
        if key in pngs:
            lines += [f'![{alt}]({pngs[key]})', '']
    if data['runningWithData']:
        lines += [f'运行中有数据的 {data["runningWithData"]} 台 VM 中，{data["cpuBelow20"]} 台 CPU 平均利用率低于 20%，{data["memoryBelow20"]} 台内存平均利用率低于 20%。'
                  '规则要求只减少利用率低的那一项，另一项容量保留。', '']
    for key, alt in [('status', 'VM 建议结果'), ('subscriptions', '各订阅每月预计节省')]:
        if key in pngs:
            lines += [f'![{alt}]({pngs[key]})', '']
    lines += summary_section(rows, total, cfg.vm_detail_max_rows)
    cc = data.get('commitmentCoverage') or {}
    if cc.get('read') and cc.get('coveredCount'):
        lines += ['', f'其中 {cc["coveredCount"]} 台建议调整的机器在 {int(cc["month"][5:])} 月有一半以上费用由 RI/SP 覆盖，涉及预计节省 '
                      f'USD {cc["coveredSavingsUSD"]:,.2f}/月。这些机器降配后，预留或节省计划的额度会转去覆盖其他机器或闲置，不会直接减少费用；'
                      f'需先交换预留或调整覆盖范围，节省在此之后才能实现。逐台情况见 Excel“{DETAIL_SHEET}”的备注列。']
    elif cc.get('read') and cc.get('month'):
        lines += ['', f'按 {int(cc["month"][5:])} 月分摊费用，建议调整的机器都没有被 RI/SP 覆盖，降配后的节省可以直接体现在账单上。']
    elif cc.get('attempted') and not cc.get('read'):
        lines += ['', '未能读取按资源的计价方式，无法判断哪些机器已被 RI/SP 覆盖；如已购买 RI/SP，实际节省可能低于预计。']
    if narrative.get('crossCheckParagraph'):
        lines += ['', narrative['crossCheckParagraph']]
    lines += ['', '## 3. 使用的规则', '',
              '- CPU和内存独立判断：准备减少的那项平均利用率<20%；另一项容量保留，每项最多减半，同代保留所需功能。',
              '- 目标VM的Cached/Uncached IOPS和吞吐平均消耗率各≤50%，不相加，不查单盘IO，也不使用Max/P99作门槛。',
              '- Batch API日平均按Count加权；要有7天有效数据和足够采样。缺指标、报价或规格信息时不推荐调整。',
              '- 按机器所在订阅、区域和可用区检查目标规格是否可部署。',
              '- 目标价格要更低，磁盘/网卡数量及架构、缓存、临时盘、安全配置要满足；同池运行节点有共同目标才一起调整。',
              '- 购买RI/SP前需确认长期用量和期限；SP不能随时取消。', '',
              '[Azure VM官网定价入口](https://azure.microsoft.com/en-us/pricing/details/virtual-machines/linux/)。只算VM基础计算费，不含磁盘、网络或付费软件。', '',
              '## 4. 订阅覆盖和其他范围', '',
              '已确认没有VM/VMSS实例的订阅：' + ('、'.join(without) if without else '无') + '。它们可能有其他服务费用，本报告不计算那些费用。', '']
    if unverified:
        lines += [unverified_text + '具体订阅见Excel“订阅范围”页。', '']
    if fallback:
        subs_text = '、'.join(sorted({f'`{u["subscriptionName"]}`' for u in fallback}))
        lines += [f'{subs_text}共{len(fallback)}台机器用批量接口读取指标时被拒绝（HTTP403），已改为逐台读取同一观察期的指标，结果已纳入分析。', '']
    lines += ['每个订阅的机器数量、建议调整数量和节省见 2.1 和Excel“订阅范围”页。缺数据/权限及目标不满足的原因汇总见 2.2，逐台原因见逐台明细的备注。', '',
              '报告仅覆盖VM规格与费用分析。闲置资源清理、其他Azure服务和更全面的架构评估不在此范围。', '',
              '## 官方依据', '',
              '- [Batch API](https://learn.microsoft.com/en-us/rest/api/monitor/metrics-batch/batch?view=rest-monitor-2023-10-01)',
              '- [Average与Count聚合](https://learn.microsoft.com/en-us/azure/azure-monitor/metrics/metrics-aggregation-explained)',
              '- [VM级磁盘消耗指标](https://learn.microsoft.com/en-us/azure/virtual-machines/disks-metrics)',
              '- [SKU能力与限制](https://learn.microsoft.com/en-us/rest/api/compute/resource-skus/list)']
    return format_and_check('\n'.join(lines) + '\n')
