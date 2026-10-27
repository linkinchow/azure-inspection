"""01 月度成本巡检报告: Excel (formulas verified) and Markdown.

Numbers, tables and coverage come from the data. Business explanations come from the narrative file
(narrative/01.json, written by the agent); where it is missing a plain default is generated from the data.
"""
from __future__ import annotations

import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties

from ..core.charts import clear as clear_charts
from ..core.excel import MONEY, verify
from ..core.fmt_md import format_and_check
from ..core.io import read_json, write_json
from . import charts as cost_charts
from .orphan import label as orphan_label
from . import commitments as cm
from . import views as sr

PERCENT = '0.0%;(0.0%);"-"'


def amount(x):
    return '未取得' if x is None else f'{x:,.2f}'


def percent(x):
    return '不适用' if x is None else f'{x:+.2%}'


def esc(x):
    return str(x).replace('|', '/').replace('\n', ' ')


def row(values):
    return '|' + '|'.join(esc(v) for v in values) + '|'


def labels_for(period):
    return {'prev': period.previous.short_label, 'cur': period.current.short_label, 'next': period.following.short_label,
            'prevDays': period.previous.days, 'curDays': period.current.days, 'nextDays': period.following.days,
            'nextMonth': period.following.month, 'curLabel': period.current.label, 'prevLabel': period.previous.label}


def cn_date(value):
    """'2026-10-14...' -> '10月14日'."""
    v = str(value or '')
    return f'{int(v[5:7])}月{int(v[8:10])}日' if len(v) >= 10 else '未取得'


def cn_span(start, end):
    a, b = str(start or ''), str(end or '')
    if len(a) < 10 or len(b) < 10:
        return '未取得'
    return f'{cn_date(a)}至{int(b[8:10])}日' if a[:7] == b[:7] else f'{cn_date(a)}至{cn_date(b)}'


def two_months(period):
    p, c = period.previous, period.current
    return f'{c.year}年{p.month}月与{c.month}月' if p.year == c.year else f'{p.label}与{c.label}'


def kind_labels(labels):
    Lp, Lc = labels['prev'], labels['cur']
    return {'curActual': f'{Lc}实际费用', 'prevActual': f'{Lp}实际费用', 'curAmortized': f'{Lc}分摊费用', 'curResources': f'{Lc}资源级费用',
            'curPricingModel': f'{Lc}计价方式', 'curResourceGroups': f'{Lc}资源组费用'}


def stem(period):
    return f'01_{period.current.short_label}月度成本巡检报告'


def growth_recommendation(subscription, service, narrative):
    custom = (narrative.get('growthRecommendations') or {}).get(f'{subscription}|{service}')
    if custom:
        return custom
    if service in ('Foundry Models', 'Cognitive Services', 'Azure OpenAI'):
        return '确认新模型和调用量变化是否符合业务计划；建议按订阅设置费用告警，跟踪模型调用量变化'
    if service == 'Virtual Machines':
        return '确认新增或扩容的机器、节点池是否为计划上线，纳入下月VM规格检查'
    if service.startswith('Azure Database'):
        return '确认新增数据库是否为计划上线，纳入数据库用量检查'
    if service == 'Log Analytics':
        return '确认日志采集范围和保留需求，见日志检查项'
    if service == 'Storage':
        return '确认数据增长和保留期限，见存储检查项'
    return '确认变化是否符合业务计划'


GRAINS = {'monthly': '月度', 'quarterly': '季度', 'annually': '年度', 'billingmonth': '账单月', 'billingquarter': '账单季度', 'billingannual': '账单年度'}


def budget_text(model, labels):
    b = model.get('budgets')
    if not b:
        return '未查询预算'
    queried, found = b.get('subscriptionsQueried', 0), b.get('budgets', [])
    with_budget = {x.get('subscriptionName') for x in found}
    text = f'对{queried}个有费用订阅查询预算：{queried - len(with_budget)}个无预算记录'
    if found:
        text += '，' + '、'.join(f'{x.get("subscriptionName")}有{GRAINS.get(str((x.get("properties") or {}).get("timeGrain", "")).lower(), "")}USD {(x.get("properties") or {}).get("amount", 0):,.0f}预算'
                                for x in found)
    rest = len(model['subscriptions']) - queried
    if rest > 0:
        text += f'；其余{rest}个订阅未查询预算'
    return text + f'。年度累计支出不作为{labels["cur"]}超预算证明'


def unreadable(model):
    """Subscriptions whose cost could not be read, grouped by reason (never counted as zero)."""
    out = {}
    for s in model['subscriptions']:
        if s['cur'] is None or s['prev'] is None:
            reason = '订阅未启用' if s.get('state') not in (None, 'Enabled') else '费用查询失败或无权限'
            out.setdefault(reason, []).append(s['name'])
    return out


def unreadable_text(model):
    groups = unreadable(model)
    if not groups:
        return ''
    parts = []
    for reason, names in groups.items():
        names, prefix = sorted(names), os.path.commonprefix(names)
        shown = f'{prefix}*' if len(names) > 3 and len(prefix) >= 4 else '、'.join(names[:3]) + (f' 等 {len(names)} 个' if len(names) > 3 else '')
        parts.append(f'{reason}（{shown}）')
    return f'另有{sum(map(len, groups.values()))}个订阅的费用无法读取：' + '；'.join(parts)


TERMS = {'P1Y': '1 年', 'P3Y': '3 年', 'P5Y': '5 年'}


def commitment_spec(c):
    if c.get('sku') or c['kind'] != '节省计划':
        return c.get('sku')
    amount = c.get('commitment')
    if not isinstance(amount, (int, float)):
        return None
    grain = '每小时' if c.get('grain') == 'Hourly' else f"每{c.get('grain') or ''}"
    return f"{grain}承诺 {c.get('currency') or ''} {amount:,.2f}".replace('  ', ' ')


def commitment_text(cs, labels):
    """Plain sentence on RI/SP coverage, existing orders and Azure's recommendations; None without data."""
    if not cs:
        return None
    parts = []
    total = sum(cs['byModel'].values())
    covered = sum(v for k, v in cs['byModel'].items() if k in cm.COMMITTED)
    if cs['pricingRead'] and total > 0:
        parts.append(f'{labels["cur"]}分摊费用中 {covered / total:.1%} 由 RI/SP 覆盖' if covered > 0
                     else f'{labels["cur"]}费用全部按需计费，没有使用 RI/SP')
    if cs['reservationsRead'] or cs['savingsPlansRead']:
        items = cs['commitments']
        low = sum(1 for c in items if any('未用满' in n for n in c['notes']))
        soon = sum(1 for c in items if any('到期' in n for n in c['notes']))
        parts.append(f'现有 {sum(1 for c in items if c["kind"] == "预留实例")} 个预留、{sum(1 for c in items if c["kind"] == "节省计划")} 个节省计划'
                     + (f'，其中 {low} 个近 30 天未用满' if low else '') + (f'，{soon} 个即将到期' if soon else ''))
    else:
        parts.append('现有 RI/SP 订单读取权限不足，未能核对使用率和到期时间')
    recs, totals = cs['recommendations'], cs.get('recommendationTotals') or {}
    if recs:
        ri, sp = totals.get('预留实例|P1Y', 0), totals.get('节省计划|P1Y', 0)
        options = [f'购买 1 年期预留实例约节省 USD {ri:,.2f}' if ri else None, f'购买 1 年期节省计划约节省 USD {sp:,.2f}' if sp else None]
        parts.append('Azure 基于过去 30 天用量估算，' + '，'.join(x for x in options if x) + '（每 30 天，按当前用量、未计入降配；'
                     '预留实例和节省计划覆盖同一部分用量，只能选其一，3 年期方案见 Excel）' if any(options) else
                     f'Azure 给出 {len(recs)} 项 3 年期购买建议，见 Excel')
    elif cs['recommendationsRead']:
        parts.append('Azure 没有给出购买建议')
    return '；'.join(parts)


def action_items(ws, model, sup, labels, narrative, period, cs=None):
    Lp, Lc = labels['prev'], labels['cur']
    pairs = {(p['subscriptionName'], p['serviceName']): p for p in sup['drivers']['pairs']}
    items = []
    for r in sorted([r for r in model['subscriptionServiceDrivers'] if r['delta'] > 0], key=lambda r: -r['delta'])[:4]:
        p = pairs.get((r['subscriptionName'], r['serviceName'])) or {'subscriptionName': r['subscriptionName'], 'serviceName': r['serviceName']}
        items.append({'topic': '费用增长', 'scope': r['subscriptionName'] + ' / ' + r['serviceName'],
                      'finding': f'{Lc}费用USD {amount(r["cur"])}，较{Lp}增加USD {amount(r["delta"])}；日均变化{percent(r["dailyChangeRate"])}。' + sr.driver_text(p, narrative),
                      'recommendation': growth_recommendation(r['subscriptionName'], r['serviceName'], narrative), 'saving': None, 'owner': '业务负责人/财务', 'status': '待业务确认',
                      'evidencePeriod': f'{two_months(period)}实际费用'})
    vm_summary = ws.month_dir('findings') / 'vm' / 'summary.json'
    if vm_summary.exists():
        s = read_json(vm_summary)
        window = s.get('observationWindow') or {}
        span = cn_span(window.get('start'), window.get('end'))
        items.append({'topic': 'VM规格', 'scope': '全部订阅中满足平均值规则的运行VM', 'finding': f'{span}检查建议调整{s["recommendedVMCount"]}台机器',
                      'recommendation': '按逐台建议先小范围调整，确认关键任务和同池应用资源要求；费用估算仅用于后续决策',
                      'saving': s['monthlyPotentialSavingsUSD'], 'owner': '云平台/业务负责人', 'status': '待审批，尚未实施',
                      'evidencePeriod': f'{span}用量与当前规格；不是{Lc}已实现节省'})
    items.append({'topic': 'RI/SP', 'scope': '长期稳定计算、数据库、缓存及可预留存储费用',
                  'finding': commitment_text(cs, labels) or '已有按服务列表价对比；是否有闲置承诺需结合现有订单及使用记录核对',
                  'recommendation': '先确定调整后的长期用量，再比较适用的1年/3年RI和SP；不要按全年满配估算直接采购',
                  'saving': None, 'owner': '采购/财务/云平台', 'status': '待核对现有订单和业务计划', 'evidencePeriod': f'当前采购建议，非{Lc}采购结果'})
    rr = model.get('resourceReview') or {}
    waste = (rr.get('wasteSummary') or {}).get('classificationCounts', {})
    snap = cn_date(rr.get('inventorySnapshotAt'))
    items.append({'topic': '闲置资源', 'scope': '未挂载磁盘、未关联IP、旧快照和停用环境',
                  'finding': (f'{snap}资源清单发现{waste.get("unattachedDiskCandidate", 0)}块未挂载磁盘、{waste.get("unassociatedPublicIPCandidate", 0)}个未关联公网IP、'
                              f'{waste.get("unassociatedNICCandidate", 0)}个未关联网卡及{waste.get("snapshotRetentionReview", 0)}个待复核快照'
                              + (f'，另有{sum(v for k, v in waste.items() if k.startswith("orphan:"))}个其他孤立资源（如空的负载均衡、未关联的 NSG、空资源组）'
                                 if any(k.startswith('orphan:') for k in waste) else '') + f'；不代表{Lc}历史状态')
                  if rr else '资源是否可删除需核对用途、引用和保留要求',
                  'recommendation': '请资源负责人确认保留或退役；批准前不删除，删除前确认备份和恢复要求',
                  'saving': None, 'owner': '资源负责人', 'status': '待用途确认', 'evidencePeriod': '当前配置检查'})
    for a in narrative.get('actions') or []:
        items.append({'topic': a['topic'], 'scope': a['scope'], 'finding': a['finding'], 'recommendation': a['recommendation'], 'saving': a.get('saving'),
                      'owner': a.get('owner', '客户负责人'), 'status': a.get('status', '待业务确认'), 'evidencePeriod': a.get('evidencePeriod', f'{Lc}用量')})
    gaps = []
    if unreadable(model):
        gaps.append(unreadable_text(model).replace('另有', ''))
    if (cs and not (cs['reservationsRead'] or cs['savingsPlansRead'])) or \
            (not cs and any(c.get('status') == 'not-verified' for c in rr.get('commitmentStatus', []))):
        gaps.append('RI/SP订单读取权限不足')
    gaps.append(f'未取得{Lc}月度预算目标和资源负责人')
    items.append({'topic': '数据与权限', 'scope': '无法读取的订阅、RI/SP订单、预算和资源负责人',
                  'finding': '；'.join(gaps) + '。读取失败或缺业务信息的项目不按零费用、零用量处理',
                  'recommendation': '按覆盖说明补充只读数据或由客户提供导出；确认预算和责任人后形成决策',
                  'saving': None, 'owner': '客户管理员/财务', 'status': '待补充', 'evidencePeriod': '报告数据覆盖说明'})
    if rr:
        cov, summ = rr.get('coverage', {}), rr.get('summary', {})
        items.append({'topic': f'{Lc}VM使用', 'scope': f'仍存在且可取得{Lc}指标的机器',
                      'finding': f'{cov.get("completeMonthCPUMemoryInstances", 0)}台有完整{labels["curDays"]}个有效日；其中{summ.get("lowCPUFullMonthCount", 0)}台CPU平均低于20%，'
                                 f'{summ.get("lowMemoryUsedFullMonthCount", 0)}台内存平均低于20%',
                      'recommendation': f'优先确认低利用率机器的业务任务与历史规格，再评估减少容量；没有历史规格证据时不直接计算{Lc}降配节省',
                      'saving': None, 'owner': '云平台/业务负责人', 'status': '待容量评审', 'evidencePeriod': f'{Lc}历史用量，资源范围为当前仍存在的机器'})
    return items


def build(ws, cfg, period, narrative=None, log=print):
    from .model import build as build_model
    narrative = narrative or {}
    labels = labels_for(period)
    Lp, Lc, Ln = labels['prev'], labels['cur'], labels['next']
    dp, dc = labels['prevDays'], labels['curDays']
    model = build_model(ws, period)
    sup = sr.load(ws)
    fdir, out = ws.month_dir('findings') / 'cost', ws.month_dir('reports')
    tot = model['comparableTotals']
    tot['amortizedCur'] = model['totals']['amortizedCur']
    raw_commitments = cm.load(ws)
    cs = cm.summary(raw_commitments) if raw_commitments else None
    actions = action_items(ws, model, sup, labels, narrative, period, cs)
    write_json(fdir / 'model.json', model)
    write_json(fdir / 'actions.json', actions)
    wb, expected = Workbook(), {}
    wb.remove(wb.active)

    def sheet(name, title, note, columns):
        s = wb.create_sheet(name)
        s.append([title])
        s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
        s.append([note])
        s.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
        s.append(columns)
        s.freeze_panes = 'A4'
        return s

    def formula(s, r, c, expression, value):
        s.cell(r, c, expression)
        expected[f'{s.title}!{L(c)}{r}'] = value

    n_comp = model['totals']['comparableSubscriptionCount']
    cover = (f'费用来自Azure Cost Management实际费用，不是官网列表价；优化建议公开报价另行注明为列表价。{Lc}{dc}天，{Lp}{dp}天，日均变化另列。'
             f'两月对比仅使用两个月均读取成功的{n_comp}个订阅；读取失败的订阅不记为0。')
    fc = sup['forecast']['overall']['next']
    overview = sheet('月度总览', f'{cfg.name}：{labels["curLabel"]}成本巡检（对比{Lp}）', cover, ['项目', '结果', '说明'])
    for r_ in [['订阅数', n_comp, f'{cfg.name}纳入巡检且两个月费用都能读取的订阅' + (f'；{unreadable_text(model)}' if unreadable(model) else '')],
               [f'{Lc}实际费用USD', None, f'完整{Lc}查询结果；未核对最终发票税费/信用及后续调整'], [f'{Lp}实际费用USD', None, f'完整{Lp}查询结果'],
               ['费用增减USD', None, f'{Lc}减{Lp}'], ['费用变化率', None, f'与{Lp}费用相比'], [f'{Lc}日均USD', None, f'{Lc}总费用/{dc}'],
               [f'{Lp}日均USD', None, f'{Lp}总费用/{dp}'], ['日均变化率', None, '排除两个月天数差异后比较'],
               [f'{Lc}分摊费用USD', tot['amortizedCur'], 'RI/SP预付费用分配到使用期间；不与实际费用相加'],
               ['数据查询时间', model['scope'].get('queriedAt') or '见原始查询记录', '完整月份费用按查询时结果列示'],
               [f'{Ln}全月预计USD', fc['fullMonthEstimate'], f'Azure Cost Management预测：已发生日期为实际，其余为预测，不是账单；见“{Ln}费用预测”']]:
        overview.append(r_)
    sw = sheet('订阅对比', f'各订阅{Lc}与{Lp}实际费用', cover,
               ['业务组', '订阅', f'{Lp}USD', f'{Lc}USD', '增减USD', '变化率', f'{Lp}日均USD', f'{Lc}日均USD', '日均变化率', f'{Lc}分摊费用USD',
                f'{Lp}读取情况', f'{Lc}读取情况', '订阅ID'])
    for i, r in enumerate(model['subscriptions'], 4):
        sw.append([r['group'], r['name'], r['prev'], r['cur'], None, None, None, None, None, r['amortizedCur'], r['actualPrevStatus'], r['actualCurStatus'], r['id']])
        if r['delta'] is not None:
            formula(sw, i, 5, f'=D{i}-C{i}', r['delta'])
        if r['changeRate'] is not None:
            formula(sw, i, 6, f'=D{i}/C{i}-1', r['changeRate'])
        if r['prev'] is not None:
            formula(sw, i, 7, f'=C{i}/{dp}', r['prevDaily'])
        if r['cur'] is not None:
            formula(sw, i, 8, f'=D{i}/{dc}', r['curDaily'])
        if r['dailyChangeRate'] is not None:
            formula(sw, i, 9, f'=H{i}/G{i}-1', r['dailyChangeRate'])
    expressions = {p: '=' + '+'.join(f"'订阅对比'!{col}{i}" for i, r in enumerate(model['subscriptions'], 4) if r['id'] in model['comparableSubscriptionIds'])
                   for p, col in [('cur', 'D'), ('prev', 'C')]}
    formula(overview, 5, 2, expressions['cur'], tot['cur'])
    formula(overview, 6, 2, expressions['prev'], tot['prev'])
    formula(overview, 7, 2, '=B5-B6', tot['delta'])
    if tot['changeRate'] is not None:
        formula(overview, 8, 2, '=B5/B6-1', tot['changeRate'])
    formula(overview, 9, 2, f'=B5/{dc}', tot['curDaily'])
    formula(overview, 10, 2, f'=B6/{dp}', tot['prevDaily'])
    if tot['dailyChangeRate'] is not None:
        formula(overview, 11, 2, '=B9/B10-1', tot['dailyChangeRate'])

    def comparison(name, title, rows):
        s = sheet(name, title, cover, ['项目', f'{Lp}USD', f'{Lc}USD', '增减USD', '变化率', f'{Lp}日均USD', f'{Lc}日均USD', '日均变化率'])
        for i, r in enumerate(rows, 4):
            s.append([r['name'], r['prev'], r['cur'], None, None, None, None, None])
            formula(s, i, 4, f'=C{i}-B{i}', r['delta'])
            if r['changeRate'] is not None:
                formula(s, i, 5, f'=C{i}/B{i}-1', r['changeRate'])
            formula(s, i, 6, f'=B{i}/{dp}', r['prevDaily'])
            formula(s, i, 7, f'=C{i}/{dc}', r['curDaily'])
            if r['dailyChangeRate'] is not None:
                formula(s, i, 8, f'=G{i}/F{i}-1', r['dailyChangeRate'])
        return s
    comparison('业务组对比', f'业务组：{Lc}与{Lp}实际费用', model['groups'])
    comparison('服务对比', f'Azure服务：{Lc}与{Lp}实际费用', model['services'])
    driver = sheet('增长与下降来源', '订阅×服务：费用变化来源', '费用变化能定位到服务和资源；是否由业务量、数量或价格变化引起，需结合业务证据确认。',
                   ['业务组', '订阅', '服务', f'{Lp}USD', f'{Lc}USD', '增减USD', '变化率', '日均变化率'])
    for i, r in enumerate(model['subscriptionServiceDrivers'], 4):
        driver.append([r['group'], r['subscriptionName'], r['serviceName'], r['prev'], r['cur'], None, r['changeRate'], r['dailyChangeRate']])
        formula(driver, i, 6, f'=E{i}-D{i}', r['delta'])
    rg = sheet('资源组对比', f'资源组：{Lc}与{Lp}实际费用', '未归属资源组的费用单独列示，不猜测资源归属。', ['订阅', '资源组', f'{Lp}USD', f'{Lc}USD', '增减USD', '变化率'])
    for i, r in enumerate(model['resourceGroups'], 4):
        rg.append([r['subscriptionName'], r['resourceGroup'], r['prev'], r['cur'], None, r['changeRate']])
        formula(rg, i, 5, f'=D{i}-C{i}', r['delta'])
    res = sheet(f'{Lc}资源费用', f'{Lc}资源级费用', '部分购买、服务级或未分配费用没有资源ID；完整费用以订阅/服务总额为准。', ['订阅', '资源组', '资源', '服务', f'{Lc}USD', '资源ID'])
    for r in model['resources']:
        res.append([r['subscriptionName'], r['resourceGroup'], r['resourceName'], r['serviceName'], r['cost'], r['resourceId']])
    day = sheet('每日费用', f'每日费用：{Lp}与{Lc}', '按月内日期排列；只有一个月有的日期不补成0。', ['日期号', f'{Lp}USD', f'{Lc}USD'])
    for r in model['daily']:
        day.append([r['day'], r['prev'], r['cur']])
    am = sheet('分摊与计价方式', f'{Lc}分摊费用及计价方式', '实际费用用于账单支出，分摊费用用于将承诺费用分配到期间；计价方式金额不等于优惠覆盖率。', ['类别', '项目', '费用USD'])
    for r in model['amortizedServices']:
        am.append(['分摊费用/服务', r['serviceName'], r['cost']])
    for r in model['pricingModels']:
        am.append(['实际费用/计价方式', r['pricingModel'], r['cost']])
    if cs:
        for k, v in sorted(cs['byModel'].items(), key=lambda kv: -kv[1]):
            am.append(['分摊费用/计价方式（按资源汇总）', k, v])
        ri = sheet('RI与SP', '现有预留实例（RI）和节省计划（SP）', '使用率来自 Azure；近 30 天低于 80% 视为未用满，90 天内到期单独提示。读不到订单时不代表没有购买。',
                   ['类别', '名称', '规格或承诺', '范围', '期限', '近7天使用率%', '近30天使用率%', '到期日期', '提示'])
        for c in cs['commitments']:
            spec = commitment_spec(c)
            u = c.get('utilization') or {}
            ri.append([c['kind'], c.get('name'), spec, c.get('scope'), TERMS.get(c.get('term'), c.get('term')), u.get('days7'), u.get('days30'), c.get('expiryDate'),
                       '；'.join(c['notes']) or '正常'])
        if not cs['commitments']:
            ri.append(['说明', '预留订单' + ('读取成功，没有预留' if cs['reservationsRead'] else '读取权限不足，未能核对'), None, None, None, None, None, None,
                       '节省计划' + ('读取成功，没有节省计划' if cs['savingsPlansRead'] else '读取权限不足，未能核对')])
        rs = sheet('Azure购买建议', 'Azure 基于过去 30 天用量的 RI/SP 购买建议', '金额为 Azure 估算的 30 天费用和节省（约等于每月），按当前用量计算，未考虑降配；'
                   '先降配并稳定后再按新的用量采购。同一订阅的 1 年和 3 年、预留实例和节省计划是互相替代的方案，节省不能相加。',
                   ['类型', '订阅', '资源类型', '规格', '区域', '期限', '建议数量或每小时承诺', '30天按需费用USD', '购买后费用USD', '估算节省USD（30天）'])
        for r in cs['recommendations']:
            rs.append([r['kind'], r['subscription'], r.get('resourceType'), r.get('sku'), r.get('region'), TERMS.get(r.get('term'), r.get('term')), r.get('quantity'),
                       r.get('costWithout'), r.get('costWith'), r.get('saving')])
    aw = sheet('建议与跟踪', '建议清单和执行跟踪', '预计节省不是已实现节省；当前检查建议与历史费用分别标注。',
               ['类别', '范围', '发现', '建议怎么做', '预计每月节省USD', '负责人', '状态', '证据时间', '客户决定', '实施日期', '实际节省USD', '业务验收结果'])
    for a in actions:
        aw.append([a['topic'], a['scope'], a['finding'], a['recommendation'], a['saving'], a['owner'], a['status'], a['evidencePeriod'], None, None, None, None])
    extra_formats = sr.add_sheets(sheet, formula, sup, model, labels, narrative, cfg)
    data = cost_charts.chart_data(model, sup, period)
    clear_charts(out / 'charts', '01')
    pngs = cost_charts.make_pngs(data, labels, out / 'charts')
    cost_charts.add_excel_sheet(sheet, data, labels)
    rr = model.get('resourceReview') or {}
    if rr.get('vmMetrics'):
        u = sheet(f'{Lc}VM使用情况', f'{Lc}历史VM使用情况', f'以当前仍存在且{Lc}底前创建的实例读取历史指标；不是{Lc}全部历史资源清单，当前规格不视为{Lc}已核实规格。',
                  ['订阅', '资源组', 'VM', '当前SKU', f'{Lc}CPU平均%', f'{Lc}内存平均%', '有效日', '说明'])
        for r in rr['vmMetrics']:
            days = r.get('validDays') or 0
            note = (f'{Lc}之后创建，不属于{Lc}观察范围' if not r.get('eligibleForMonth', True) else f'有{dc}个有效日；降配前需核对{Lc}规格和业务需求' if days == dc
                    else f'只有{days}个有效日，数据不足整月' if days > 0 else f'缺少{Lc}指标，不能判断用量')
            if any('403' in str(m.get('error', '')) for m in r.get('metrics', {}).values()):
                note = '指标读取权限不足（HTTP403），暂不判断'
            u.append([r.get('subscriptionName'), r.get('resourceGroup'), r.get('name'), r.get('currentSKU'), r.get('cpuAverage'), r.get('memoryAverage'), r.get('validDays'), note])
    if rr.get('wasteFindings'):
        waste_sheet = sheet('当前闲置资源待确认', '当前资源清单：待确认用途的资源', f'这是当前资源检查，不是{Lc}历史闲置或节省证明；不自动删除。',
                            ['类别', '订阅', '资源组', '资源', '原因', '快照时间', '资源ID'])
        cat_labels = {'unattachedDiskCandidate': '未挂载磁盘，待确认用途', 'unassociatedPublicIPCandidate': '未关联公网IP，待确认用途',
                      'unassociatedNICCandidate': '未关联网卡，待确认用途', 'snapshotRetentionReview': '快照保留期复核'}
        for r in rr['wasteFindings']:
            category = r.get('category')
            reason = ('需确认快照用途和保留时间，不按年龄直接删除' if category == 'snapshotRetentionReview'
                      else '当前配置下没有关联或使用（孤立资源查询）；' + ('闲置时仍收费，' if r.get('charged') else '') + '需要负责人确认后再删除'
                      if str(category).startswith('orphan:') else '当前配置无常见关联；需要负责人确认用途、备份和保留要求')
            waste_sheet.append([cat_labels.get(category) or orphan_label(category) or category, r.get('subscriptionName'), r.get('resourceGroup'), r.get('name'), reason, rr.get('inventorySnapshotAt'), r.get('resourceId')])
    budget_sheet = sheet('预算检查', f'现有预算及{Lc}可判断范围', f'只列实际读取到的预算；年度累计支出不是{Lc}单月支出，不能按年度预算直接判断{Lc}超预算。',
                         ['订阅', '预算名称', '金额', '币种', '周期', '开始日期', '结束日期', '当前累计支出', '说明'])
    budgets = (model.get('budgets') or {}).get('budgets', [])
    for b in budgets:
        props = b.get('properties', {})
        p_, spend = props.get('timePeriod', {}), props.get('currentSpend', {})
        monthly = str(props.get('timeGrain', '')).lower() == 'monthly'
        budget_sheet.append([b.get('subscriptionName'), b.get('name'), props.get('amount'), spend.get('unit', 'USD'), props.get('timeGrain'), p_.get('startDate'), p_.get('endDate'),
                             spend.get('amount'), '月度预算' if monthly else f'非月度预算；未取得{Lc}月度目标，不能判断{Lc}是否超预算'])
    coverage = sheet('数据覆盖与说明', '数据覆盖和未完成检查', '读取失败不按0处理；没有实际执行记录时不填写已实现节省。', ['项目', '说明'])
    for kind, c in model['scope']['coverage'].items():
        coverage.append([kind_labels(labels).get(kind, kind), f'{c["success"]}个订阅读取成功' + (f'；{len(c["failed"])}个订阅读取失败' if c['failed'] else '')])
    if unreadable(model):
        coverage.append(['费用覆盖限制', unreadable_text(model) + f'；未取得的费用不按0处理。所有环比使用{n_comp}个订阅'])
    disabled = unreadable(model).get('订阅未启用') or []
    for r_ in [['历史费用', f'Cost Management ActualCost；{Lc}与{Lp}完整自然月，不使用月底预测替代'], ['月份天数', f'{Lc}{dc}天、{Lp}{dp}天，同时比较日均费用'],
               ['预算', '如未取得预算金额，只报告费用变化，不判断是否超预算'], ['RI/SP现有订单', '是否买多/用满需现有订单和使用记录；价格对比不证明已购买或已节省'],
               ['优化报价', f'当前公开列表价用于后续决策，不是{Lc}实际账单费用'], ['当前资源状态', f'当前盘/VM/资源配置仅列为后续建议，不当作{Lc}状态'],
               ['实际效果', '尚未提供执行及验收记录，不填已实现节省'], ['WAF参考', 'https://learn.microsoft.com/en-us/azure/well-architected/cost-optimization/checklist']]:
        coverage.append(r_)
    coverage.append(['预算读取', budget_text(model, labels)])
    if cs:
        coverage.append(['RI/SP 覆盖（按资源）', f'{cs["pricingRead"]} 个有费用的订阅读取成功（{Lc}分摊费用，按资源和计价方式）'])
        coverage.append(['RI/SP 现有订单', ('预留：读取成功；' if cs['reservationsRead'] else '预留：读取权限不足，未能核对使用率和到期；')
                         + ('节省计划：读取成功' if cs['savingsPlansRead'] else '节省计划：读取权限不足，未能核对')])
        coverage.append(['Azure 购买建议', f'{cs["recommendationsRead"]}/{cs["recommendationsTried"]} 次查询成功，共 {len(cs["recommendations"])} 条建议'])
    else:
        for c in rr.get('commitmentStatus', []):
            coverage.append(['RI/SP订单读取', c.get('kind', '') + ('：读取权限不足，未能核对现有订单和使用率' if c.get('status') == 'not-verified' else '：' + c.get('status', '未核对'))])
    for r_ in sr.coverage_rows(sup, model, labels):
        coverage.append(r_)
    if disabled:
        coverage.append(['未启用订阅', f'{len(disabled)}个未启用订阅在资源查询中不可见，其数据库、存储和日志未检查'])
    if rr:
        rc = rr.get('coverage', {})
        coverage.append([f'{Lc}VM历史数据', f'{rc.get("eligibleMonthInstances", 0)}台在{Lc}已存在，{rc.get("completeMonthCPUMemoryInstances", 0)}台有{dc}个有效日，'
                                            f'{rc.get("partialCPUMemoryInstances", 0)}台部分数据，{rc.get("missingBothInstances", 0)}台缺CPU/内存；'
                                            f'另{rc.get("excludedCreatedAfterMonth", 0)}台在{Lc}之后创建，不计入{Lc}用量'])
    for s in wb:
        s.sheet_view.showGridLines = False
        s.auto_filter.ref = f'A3:{L(s.max_column)}{s.max_row}'
        s.print_title_rows = '1:3'
        s.page_setup.orientation = 'landscape'
        s.page_setup.fitToWidth = 1
        s.row_dimensions[1].height, s.row_dimensions[2].height, s.row_dimensions[3].height = 28, 52, 48
        for row_ in s:
            for c in row_:
                c.font = Font(name='Arial', size=10)
                c.alignment = Alignment(wrap_text=True, vertical='top')
                if isinstance(c.value, (int, float)) or c.data_type == 'f':
                    c.number_format = MONEY
                if c.data_type == 'f':
                    c.font = Font(name='Arial', size=10, color='008000' if '!' in c.value else '000000')
                if c.row == 1:
                    c.fill = PatternFill('solid', fgColor='17365D')
                    c.font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
                elif c.row == 3:
                    c.fill = PatternFill('solid', fgColor='DCE6F1')
                    c.font = Font(name='Arial', size=10, bold=True)
                elif c.row > 3 and c.row % 2 == 0:
                    c.fill = PatternFill('solid', fgColor='F5F8FC')
                if isinstance(c.value, str) and c.value.startswith('https://'):
                    c.hyperlink = c.value
        for col in range(1, s.max_column + 1):
            s.column_dimensions[L(col)].width = 24
        s.column_dimensions['A'].width = 38
    for s in [overview, coverage]:
        s.column_dimensions['B'].width = 85
    overview.column_dimensions['C'].width = 110
    for cols, s in [([6, 9], sw), ([7, 8], driver)]:
        for col in cols:
            for r_ in s.iter_rows(min_row=4, min_col=col, max_col=col):
                r_[0].number_format = PERCENT
    for name in ['业务组对比', '服务对比']:
        for col in [5, 8]:
            for r_ in wb[name].iter_rows(min_row=4, min_col=col, max_col=col):
                r_[0].number_format = PERCENT
    for r_ in [8, 11]:
        overview.cell(r_, 2).number_format = PERCENT
    for s, cols in extra_formats:
        for col in cols:
            for r_ in s.iter_rows(min_row=4, min_col=col, max_col=col):
                r_[0].number_format = PERCENT
    for row_ in wb['数据库与缓存用量'].iter_rows(min_row=4, min_col=12, max_col=12):
        row_[0].number_format = '0'
    for name, col in [('数据库与缓存用量', 13), ('变化原因拆解', 12), ('日志摄取与保留', 8), ('存储容量与层级', 12)]:
        wb[name].column_dimensions[L(col)].width = 80
    order = ['月度总览', '图表', '订阅对比', '业务组对比', '服务对比', '区域对比', '变化原因拆解', '主要计量变化', '增长与下降来源', '资源组对比', f'{Lc}资源费用', '每日费用',
             f'{Ln}费用预测', '分摊与计价方式', 'RI与SP', 'Azure购买建议', '建议与跟踪', f'{Lc}VM使用情况', '数据库与缓存用量', '存储容量与层级', '日志摄取与保留', '当前闲置资源待确认', '预算检查', '数据覆盖与说明']
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]
    aw.column_dimensions['B'].width, aw.column_dimensions['C'].width, aw.column_dimensions['D'].width, aw.column_dimensions['H'].width = 60, 85, 100, 65
    for i in range(4, aw.max_row + 1):
        aw.row_dimensions[i].height = 85
    out.mkdir(parents=True, exist_ok=True)
    path = out / f'{stem(period)}.xlsx'
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
    write_json(fdir / 'expected-formulas.json', expected)
    md = markdown(cfg, period, labels, model, actions, sup, pngs, narrative, cs)
    (out / f'{stem(period)}.md').write_text(md, encoding='utf-8')
    result = verify(path, expected, fdir / 'verification.json')
    log(f'01 报告：{Lc}费用 USD {tot["cur"]:,.2f}（{n_comp}个订阅），较{Lp} {percent(tot["changeRate"])}；Excel 公式 {result["formulaCount"]} 个'
        + ('，已用 Excel 复算核对' if result['recalculatedInExcel'] else '，本机未能用 Excel 复算'))
    return {'markdown': out / f'{stem(period)}.md', 'excel': path, 'model': model}


def _day(d):
    return f'{d.year}年{d.month}月{d.day}日'


def markdown(cfg, period, labels, m, actions, sup, pngs, narrative, cs=None):
    Lp, Lc, Ln = labels['prev'], labels['cur'], labels['next']
    dp, dc, dn = labels['prevDays'], labels['curDays'], labels['nextDays']
    t, cov = m['comparableTotals'], m['scope']['coverage']
    nx = sup['forecast']['overall']['next']
    db, st, lg = sup['databases'], sup['storage'], sup['logs']
    notes = narrative.get('chartNotes') or {}
    q = m['scope'].get('queriedAt') or ''
    queried = f'{q[:4]}年{int(q[5:7])}月{int(q[8:10])}日 {q[11:16]}' if len(q) >= 16 else '见查询记录'
    n_comp = m['totals']['comparableSubscriptionCount']
    unread = unreadable_text(m)
    lines = [f'# {cfg.name} {labels["curLabel"]}月度成本巡检报告', '',
             f'**统计期间：{_day(period.current.first)}至{period.current.last.day}日；对比{_day(period.previous.first)}至{period.previous.last.day}日。'
             f'范围：{cfg.name} {n_comp}个Azure订阅' + (f'（另有{sum(map(len, unreadable(m).values()))}个订阅的费用无法读取，见第8节）' if unread else '') + '。金额单位：USD。**', '',
             '历史费用来自Azure Cost Management的实际费用查询，不是官网列表价；优化建议中的公开报价另行标注为列表价。', '',
             f'{Lp}和{Lc}费用查询完成时间：{queried}（UTC）；区域、预测、变化原因、数据库、存储和日志检查于{cn_date(db["snapshotAt"])}读取。', '',
             '## 1. 费用总览', '',
             f'- {Lc}费用（{n_comp}个订阅） **USD {amount(t["cur"])}**，{Lp} **USD {amount(t["prev"])}**，增减 **USD {amount(t["delta"])}（{percent(t["changeRate"])}）**。',
             f'- {Lc}日均 **USD {amount(t["curDaily"])}**，{Lp}日均 **USD {amount(t["prevDaily"])}**，日均变化 **{percent(t["dailyChangeRate"])}**。'
             + (f'{Lc}比{Lp}{"少" if dc < dp else "多"}{abs(dp - dc)}天，月总额与日均值一起看。' if dc != dp else ''),
             f'- 实际费用查询：{cov["curActual"]["success"]}个订阅的{Lc}费用、{cov["prevActual"]["success"]}个订阅的{Lp}费用均已读取，以上增减比较使用这{n_comp}个订阅'
             + (f'；{unread}，没有按0处理' if unread else '') + '。']
    if narrative.get('forecastSummary'):
        lines.append(f'- {narrative["forecastSummary"]}')
    elif nx['fullMonthEstimate']:
        est = nx['fullMonthEstimate']
        lines.append(f'- **{Ln}预计约USD {amount(est)}**（Azure预测）：截至查询日实际USD {amount(nx["actualToDate"])}，其余日期预测USD {amount(nx["forecastRemainder"])}。'
                     f'全月日均约USD {amount(est / dn)}，比{Lc}日均{"高" if est / dn > t["curDaily"] else "低"}{abs(est / dn / t["curDaily"] - 1):.2%}；预测需在{Ln}结束后核对。')
    lines += ['- 费用按查询时结果列示；税费、信用额度和最终发票调整需与财务账单核对。', '']
    if 'daily' in pngs:
        lines += [f'![{Lp}与{Lc}每日费用]({pngs["daily"]})', '',
                  notes.get('daily') or f'{Lc}日均 USD {amount(t["curDaily"])}，比{Lp}（USD {amount(t["prevDaily"])}）{"高" if (t["dailyChangeRate"] or 0) >= 0 else "低"} {abs(t["dailyChangeRate"] or 0):.2%}。', '']
    if 'forecast' in pngs:
        lines += [f'![{Ln}每日费用：实际与预测]({pngs["forecast"]})', '',
                  notes.get('forecast') or f'实线为{Ln}已发生的实际费用，虚线为 Azure 预测；预测基于历史费用模型，节假日和新上线业务可能使实际偏离预测。', '']
    lines += ['### 需要优先关注的变化', '']
    highlights = narrative.get('highlights')
    if not highlights:
        highlights = []
        if dc != dp:
            highlights.append(f'**看日均，不只看月总额。**{Lc}{dc}天、{Lp}{dp}天，月总额变化 {percent(t["changeRate"])}，日均变化 {percent(t["dailyChangeRate"])}。')
        ups = sorted(m['subscriptionServiceDrivers'], key=lambda r: -r['delta'])
        if ups and ups[0]['delta'] > 0:
            r = ups[0]
            highlights.append(f'**增加最多：{r["subscriptionName"]} / {r["serviceName"]}**，USD {amount(r["delta"])}。原因见第3节。')
        if ups and ups[-1]['delta'] < 0:
            r = ups[-1]
            highlights.append(f'**减少最多：{r["subscriptionName"]} / {r["serviceName"]}**，USD {amount(r["delta"])}。')
    group = cost_charts.group_label(m)
    lines += [f'- {h}' for h in highlights] + ['', f'## 2. 哪些{"业务" if group == "业务组" else "订阅"}、服务和区域变化最大', '',
                                               f'|{group}|{Lp}费用USD|{Lc}费用USD|增减USD|变化率|日均变化率|', '|---|---:|---:|---:|---:|---:|']
    for r in cost_charts.group_rows(m):
        lines.append(row([r['name'], amount(r['prev']), amount(r['cur']), amount(r['delta']), percent(r['changeRate']), percent(r['dailyChangeRate'])]))
    if 'groups' in pngs:
        lines += ['', f'![各{group}{Lp}与{Lc}费用]({pngs["groups"]})']
    lines += ['', '### 主要服务', '', f'|服务|{Lp}费用USD|{Lc}费用USD|增减USD|日均变化率|', '|---|---:|---:|---:|---:|']
    for r in m['services'][:15]:
        lines.append(row([r['name'], amount(r['prev']), amount(r['cur']), amount(r['delta']), percent(r['dailyChangeRate'])]))
    if 'services' in pngs:
        lines += ['', f'![费用增减最大的服务]({pngs["services"]})']
    if notes.get('services'):
        lines += ['', notes['services']]
    reg = sup['region']['overall']
    lines += ['', '### 按资源所在区域', '', '这里按资源实际部署区域统计，与上面的订阅业务分组不同。', '',
              f'|区域|{Lp}费用USD|{Lc}费用USD|增减USD|{Lc}占比|日均变化率|', '|---|---:|---:|---:|---:|---:|']
    for r in reg['byLocation'][:8]:
        lines.append(row([sr.region_label(r['location'], cfg), amount(r['prev']), amount(r['cur']), amount(r['cur'] - r['prev']), f'{r["cur"] / reg["totals"]["cur"]:.1%}',
                          percent((r['cur'] / dc) / (r['prev'] / dp) - 1) if r['prev'] else '不适用']))
    lines += [''] + ([notes['regions'], ''] if notes.get('regions') else ['其余区域见Excel“区域对比”。', ''])
    if 'regions' in pngs:
        lines += [f'![各区域费用增减]({pngs["regions"]})', '']
    lines += ['## 3. 费用变化原因', '', f'下表按资源和计量比较了{Lp}与{Lc}的费用和用量，前10项增长和前5项下降均已拆分到具体资源。计量名称是Azure账单中的计费项目名称。', '',
              f'|订阅|服务|{Lp}USD|{Lc}USD|增减USD|主要原因|', '|---|---|---:|---:|---:|---|']
    for p in [p for p in sup['drivers']['pairs'] if p.get('status') == 'success']:
        lines.append(row([p['subscriptionName'], p['serviceName'], amount(p['totals']['prev']), amount(p['totals']['cur']), amount(p['totals']['delta']), sr.driver_text(p, narrative)]))
    lines += ['', '“单位费用变化”按实际费用除以计费数量计算，可能来自价格调整或折扣，账单数据无法区分。新增、停用、用量、天数和单位费用的拆分金额见Excel“变化原因拆解”，'
              '主要计量见“主要计量变化”。业务上为什么增加（如新项目上线、新模型试用）仍需业务团队确认。', '',
              '## 4. 主要资源费用', '', f'|订阅|资源组|资源|服务|{Lc}费用USD|', '|---|---|---|---|---:|']
    for r in m['resources'][:20]:
        lines.append(row([r['subscriptionName'], r['resourceGroup'], r['resourceName'], r['serviceName'], amount(r['cost'])]))
    lines += ['', '完整订阅、服务、资源组、资源和每日费用见Excel。未归属资源的费用保持未归属，不根据名称猜测。', '',
              '## 5. 资源利用率与后续优化建议', '', '### 虚拟机', '',
              f'{Lc}历史费用与当前资源建议分开。当前仍存在的机器可读取{Lc}历史指标，但不能还原{Lc}已删除的全部资源，也不能把当前规格和电源状态当作{Lc}状态。', '']
    rr = m.get('resourceReview') or {}
    if rr:
        rc, stats = rr.get('coverage', {}), rr.get('summary', {})
        lines += [f'可查询{Lc}历史指标的当前机器有{rc.get("eligibleMonthInstances", 0)}台：{rc.get("completeMonthCPUMemoryInstances", 0)}台有完整{dc}个有效日，'
                  f'{rc.get("partialCPUMemoryInstances", 0)}台只有部分日期，{rc.get("missingBothInstances", 0)}台缺CPU/内存数据；另有{rc.get("excludedCreatedAfterMonth", 0)}台在{Lc}之后才创建，不计入{Lc}用量。', '',
                  f'完整{dc}个有效日的机器中，{stats.get("lowCPUFullMonthCount", 0)}台CPU平均低于20%，{stats.get("lowMemoryUsedFullMonthCount", 0)}台内存平均低于20%。'
                  '这说明值得进一步做容量评审，不等于这些机器都能直接降配。', '']
        high = [r for r in rr.get('vmMetrics', []) if r.get('validDays') == dc and r.get('memoryAverage') is not None and r['memoryAverage'] >= 80]
        if high:
            lines += ['**高内存用量需保留余量：**' + '；'.join(f'{r["subscriptionName"]}/{r["name"]}：内存平均{r["memoryAverage"]:.2f}%' for r in high[:3])
                      + '。这表示月平均较高，不代表已确认发生性能故障；不建议直接减少内存。', '']
    lines += ['### 数据库和缓存', '', f'按{Lc}每日平均值计算；规格为{cn_date(db["snapshotAt"])}的当前配置。“CPU和内存均低于20%”只作为容量评审线索，调整前需确认峰值、高可用和副本需求。', '',
              f'|类型|数量|{Lc}费用USD|完整{dc}日|CPU和内存均低于20%|这些资源{Lc}费用USD|', '|---|---:|---:|---:|---:|---:|']
    for k, v in db['coverage'].items():
        low = [r for r in db['resources'] if r['type'].lower() == k and sr.is_low(r)]
        lines.append(row([sr.TYPE_LABELS.get(k, k), v['found'], amount(v['curCostUSD']), v.get('complete', 0), len(low), amount(sum(r['curCostUSD'] or 0 for r in low))]))
    watch = set(narrative.get('notableDatabases') or [])
    excluded = {x['name']: x['reason'] for x in narrative.get('downsizeExclusions') or []}
    notable = [r for r in db['resources'] if ((sr.is_low(r) and (r['curCostUSD'] or 0) >= 300) or r['name'] in watch) and r['name'] not in excluded]
    if notable:
        lines += ['', '费用较高的低使用率资源：', '', f'|订阅|名称|规格|{Lc}费用USD|CPU平均|内存平均|实际内存使用|有效日|', '|---|---|---|---:|---:|---:|---:|---:|']
        for r in sorted(notable, key=lambda r: -(r['curCostUSD'] or 0)):
            av = r['averages']
            gib = av.get('usedMemoryBytesGiB')
            lines.append(row([r['subscriptionName'], r['name'], sr.sku_text(r), amount(r['curCostUSD']), f'{av["cpu"]:.1f}%', f'{av["memory"]:.1f}%',
                              f'{gib:.2f} GiB' if gib is not None else '-', sr.min_days(r)]))
    lines += [''] + [f'{name}{Lc}利用率低，但{reason}，不列为降档候选。' for name, reason in excluded.items()] + ([''] if excluded else []) + [
              notes.get('databases') or 'Redis的内存百分比按整数上报，低于1%会显示为0%，因此同时列出实际内存使用量。完整清单见Excel“数据库与缓存用量”。', '',
              '### 建议清单', '', '|检查方向|发现或需要确认的内容|建议怎么做|状态|', '|---|---|---|---|']
    for a in actions:
        lines.append(row([a['topic'] + ' / ' + a['scope'], a['finding'], a['recommendation'], a['status']]))
    lines += ['', f'其中VM规格调整建议的预计节省按**官网公开列表价**计算，仅供后续调整决策，不是{Lc}实际费用减少。数据库、缓存、存储和日志的建议需先确认业务要求，暂不估算节省金额。', '']
    commitments = rr.get('commitmentStatus', [])
    text = commitment_text(cs, labels)
    if text:
        lines += ['### RI 与 SP', '', text + '。', '']
        flagged = [c for c in cs['commitments'] if c['notes']]
        if flagged:
            lines += ['|类别|名称|规格或承诺|近30天使用率|到期日期|提示|', '|---|---|---|---:|---|---|']
            for c in flagged[:10]:
                u30 = (c.get('utilization') or {}).get('days30')
                lines.append(row([c['kind'], c.get('name'), commitment_spec(c), f'{u30:.0f}%' if isinstance(u30, (int, float)) else '未取得',
                                  c.get('expiryDate') or '-', '；'.join(c['notes'])]))
            lines.append('')
        if cs['recommendations']:
            lines += ['Azure 的购买建议按当前用量计算，没有考虑降配。建议先完成降配并观察稳定，再按新的长期用量确定 RI 或 SP 的数量和期限。明细见 Excel“Azure购买建议”。', '']
    elif any(c.get('status') == 'not-verified' for c in commitments):
        lines += ['RI/SP现有订单清单读取权限不足，未能核对是否已购买、是否用满及到期时间；成本分摊查询可用，但不能代替订单使用率检查。', '']
    lines += ['## 6. 存储、备份与日志', '', '### 存储', '', f'共{st["coverage"]["found"]}个存储账户，{Lc}费用USD {amount(st["coverage"]["curCostUSD"])}。{Lc}费用USD 50以上的账户：', '',
              f'|订阅|存储账户|冗余|{Lc}底容量|{Lc}增长|层级|生命周期策略|{Lc}费用USD（含带宽）|', '|---|---|---|---:|---:|---|---|---:|']
    for a in sorted([a for a in st['accounts'] if (a['curCostUSD'] or 0) >= 50], key=lambda a: -a['curCostUSD']):
        u, lc = a.get('usedCapacity') or {}, a.get('lifecyclePolicy') or {}
        delete_only = lc.get('status') == '已配置' and all(set(x.get('actions', [])) <= {'delete'} for x in lc.get('rules', []))
        cap = '未取得' if u.get('lastGiB') is None else (f'{u["lastGiB"]:,.0f} GiB' if u['lastGiB'] >= 1 else f'{u["lastGiB"]:.2f} GiB')
        created = a.get('createdAt') or ''
        growth = (f'{u["growthPct"]:+.0f}%' if u.get('growthPct') is not None else
                  f'{Lc}{int(created[8:10])}日新建' if created.startswith(period.current.key) else '-')
        tier = 'Premium文件共享' if a.get('kind') == 'FileStorage' else ('，'.join(a.get('tierMixLastDayGiB') or {}) or a.get('accessTier'))
        lines.append(row([a['subscriptionName'], a['name'], a.get('redundancy'), cap, growth, tier, '只有删除规则' if delete_only else lc.get('status', '未取得'), amount(a['curCostUSD'])]))
    no_policy = st['coverage']['lifecycle'].get('未配置', 0)
    lines += [''] + ([narrative['storageTableNote'], ''] if narrative.get('storageTableNote') else []) + [f'- {x}' for x in (narrative.get('storageNotes') or [])] + [f'- {no_policy}个存储账户没有生命周期策略，建议按数据保留要求确认是否需要转冷或删除规则。',
                                                                                 '- 本次未检查备份保管库的备份策略；快照清单见Excel“当前闲置资源待确认”。', '',
                                                                                 '### 日志', '', f'共{lg["coverage"]["found"]}个Log Analytics工作区，{Lc}费用USD {amount(lg["coverage"]["curCostUSD"])}。{Lc}费用较高的{len([w for w in lg["workspaces"] if w.get("top3")])}个工作区：', '']
    tail = lg['period'].get('sourceWindow') or ['', '']
    no_forecast = len([s for s in sup['forecast']['subscriptions'] if s['next'].get('fullMonthEstimate') is None])
    lines += [f'|订阅|工作区|{Lc}费用USD|{Lc}可计费GB|主要数据类型|保留天数|主要来源（{cn_date(tail[0])}至{period.current.last.day}日）|', '|---|---|---:|---:|---|---:|---|']
    for w in [w for w in lg['workspaces'] if w.get('top3')]:
        rows_ = (w.get('topDataTypeBreakdown') or {}).get('rows') or []
        src = '；'.join((x.get('Category') or x.get('ContainerName') or '') + (f' {x.get("BilledGB", 0):,.0f} GB' if x.get('BilledGB', 0) >= 1 else ' 不足1 GB') for x in rows_[:2])
        top_type = (w.get('byDataType') or [{}])[0].get('DataType', '-')
        lines.append(row([w['subscriptionName'], w['name'], amount(w['curCostUSD']), f'{w["billableGB"]:,.0f}', top_type, w.get('retentionInDays'), src or '-']))
    dup = lg['coverage'].get('resourcesWithDuplicateDestinations', 0)
    lines += [''] + [f'- {x}' for x in (narrative.get('logNotes') or [])] + [
        f'- 检查了{lg["coverage"]["diagnosticSettingsChecked"]}个资源的诊断设置，' + (f'{dup}个资源有同一类日志发送到多个工作区，建议合并。' if dup else '没有发现同一类日志重复发送到多个工作区。'), '',
        '### 网络与架构', ''] + [f'- {x}' for x in (narrative.get('networkNotes') or ['网络费用变化已包含在第3节。'])] + [
        '- 架构成本评审（如节点池整合、模型部署方式、数据存放位置）需要业务和技术团队一起讨论，本次报告只提供费用和用量依据。', '',
        '## 7. 月度评审与执行跟踪', '',
        '建议业务、技术和财务一起确认主要费用变化是否符合计划，选择优先处理的项目，并补齐预算、资源负责人及采购信息。批准后的调整先小范围实施；业务验收和费用结果填写到Excel“建议与跟踪”，未执行的项目不填写已实现节省。', '',
        '## 8. 数据范围与限制', '',
        f'- {Lc}和{Lp}均使用完整自然月ActualCost查询，不以月中或月底预测代替历史费用。']
    if unread:
        lines.append(f'- {unread}；未取得其费用，不能确认为0。本文费用是{n_comp}个订阅的合计，最终费用以账单为准。')
    lines += [f'- {Ln}费用为Azure Cost Management的预测值，不是账单' + (f'；{no_forecast}个订阅没有返回预测结果（多为{Lc}无费用的订阅）。' if no_forecast else '。'),
              f'- 数据库、缓存、存储和日志的规格和配置为{cn_date(db["snapshotAt"])}读取的当前状态；使用率和用量是{Lc}历史数据。',
              f'- 日志来源明细只查询了{cn_date(tail[0])}至{period.current.last.day}日；全月摄取量来自工作区用量记录。',
              f'- 成本管理费用不等于最终发票；没有{Lc}月度预算目标时不判断是否超预算，没有执行记录时不声称已经省钱。',
              '- RI/SP订单、预算及订阅指标权限的读取情况在Excel“数据覆盖与说明”列出；缺失记录不等于没有订单或预算。',
              f'- 当前资源配置与公开列表价来自巡检时读取，单独标注，不能追溯为{Lc}部署状态或已实现收益。', '',
              '## 参考', '',
              '- [Microsoft WAF Cost Optimization](https://learn.microsoft.com/en-us/azure/well-architected/cost-optimization/checklist)',
              '- [收集和复核费用数据](https://learn.microsoft.com/en-us/azure/well-architected/cost-optimization/collect-review-cost-data)',
              '- [优化存储、保留和备份费用](https://learn.microsoft.com/en-us/azure/well-architected/cost-optimization/optimize-data-costs)',
              '- [Azure Cost Management](https://learn.microsoft.com/en-us/azure/cost-management-billing/costs/overview-cost-management)']
    text = re.sub(r'USD(?=\d)', 'USD ', '\n'.join(lines) + '\n')
    return format_and_check(text)
