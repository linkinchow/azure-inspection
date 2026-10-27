"""Independent recheck of the 01 report against raw data: totals, groupings, workbook values, supplements and Markdown.

Raises AssertionError on the first mismatch; writes findings/cost/audit.json on success.
"""
from __future__ import annotations

import math

from markdown_it import MarkdownIt
from openpyxl import load_workbook

from ..core.io import read_json, write_json
from . import charts as cost_charts
from .report import labels_for, stem


def close(a, b, what=''):
    assert a is not None and b is not None and math.isclose(a, b, abs_tol=1e-5, rel_tol=1e-9), (what, a, b)


def audit(ws, period, log=print):
    labels = labels_for(period)
    Lc, Ln = labels['cur'], labels['next']
    dp, dc = period.previous.days, period.current.days
    fdir, out = ws.month_dir('findings') / 'cost', ws.month_dir('reports')
    data, model = read_json(ws.raw('cost/cost-data.json')), read_json(fdir / 'model.json')
    assert data['complete']
    source = {(i['subscriptionId'], i['queryKind']): i for i in data['index']}
    actual_cur, actual_prev, amort = data['rows']['curActual'], data['rows']['prevActual'], data['rows'].get('curAmortized', [])
    first_c, last_c = period.current.first.isoformat(), period.current.last.isoformat()
    first_p, last_p = period.previous.first.isoformat(), period.previous.last.isoformat()
    assert all(first_c <= str(r['date'])[:10] <= last_c for r in actual_cur if r.get('date')), '本月费用含月外日期'
    assert all(first_p <= str(r['date'])[:10] <= last_p for r in actual_prev if r.get('date')), '上月费用含月外日期'
    for r in model['subscriptions']:
        for field, kind, rows in [('prev', 'prevActual', actual_prev), ('cur', 'curActual', actual_cur), ('amortizedCur', 'curAmortized', amort)]:
            q = source.get((r['id'], kind), {})
            if q.get('ok'):
                value = math.fsum(float(x['cost']) for x in rows if x['subscriptionId'] == r['id'])
                close(r[field], value, f'{r["name"]} {field}')
                if q.get('totalCost') is not None:
                    close(q['totalCost'], value, f'{r["name"]} {kind}')
            else:
                assert r[field] is None, f'{r["name"]} {field} 读取失败却有数值'
        if r['prev'] is not None and r['cur'] is not None:
            close(r['delta'], r['cur'] - r['prev'])
            close(r['prevDaily'], r['prev'] / dp)
            close(r['curDaily'], r['cur'] / dc)
            if r['prev'] > 0:
                close(r['dailyChangeRate'], (r['cur'] / dc) / (r['prev'] / dp) - 1)
    for field, rows in [('prev', actual_prev), ('cur', actual_cur), ('amortizedCur', amort)]:
        close(model['totals'][field], math.fsum(float(r['cost']) for r in rows), field)
    comparable, comparison = set(model['comparableSubscriptionIds']), model['comparableTotals']
    for name, field in [('groups', 'group'), ('services', 'serviceName')]:
        for p, rows in [('prev', actual_prev), ('cur', actual_cur)]:
            close(math.fsum(r[p] for r in model[name]), comparison[p], name)
            for row in model[name]:
                close(row[p], math.fsum(r['cost'] for r in rows if r['subscriptionId'] in comparable and (r.get(field) or '未归属') == row['name']), row['name'])
    for p in ('prev', 'cur'):
        close(math.fsum(r[p] for r in model['subscriptionServiceDrivers']), comparison[p], 'drivers')
        close(math.fsum(r[p] for r in model['resourceGroups']), comparison[p], 'resourceGroups')
    path = out / f'{stem(period)}.xlsx'
    wb = load_workbook(path, data_only=True, read_only=True)
    close(wb['月度总览']['B5'].value, comparison['cur'])
    close(wb['月度总览']['B6'].value, comparison['prev'])
    close(wb['月度总览']['B7'].value, comparison['delta'])
    sw = wb['订阅对比']
    assert sw.max_row == len(model['subscriptions']) + 3
    for i, r in enumerate(model['subscriptions'], 4):
        assert sw.cell(i, 2).value == r['name'] and sw.cell(i, 13).value == r['id'], r['name']
        for col, field in [(3, 'prev'), (4, 'cur'), (5, 'delta'), (10, 'amortizedCur')]:
            if r[field] is None:
                assert sw.cell(i, col).value is None, (r['name'], field)
            else:
                close(sw.cell(i, col).value, r[field], (r['name'], field))
    for row in wb['建议与跟踪'].iter_rows(min_row=4):
        assert row[10].value is None, '没有执行证据却填写了已实现节省'
    base = ws.raw('cost')
    region, forecast, drivers = (read_json(base / f'supplement/{n}.json') for n in ('region', 'forecast', 'drivers'))
    for n in ('resources/databases', 'resources/storage', 'resources/logs'):
        assert read_json(base / f'{n}.json')['complete'], n
    by_id = {s['id']: s for s in model['subscriptions']}
    for p in ('prev', 'cur'):
        close(math.fsum(r[p] for r in region['overall']['byLocation']), comparison[p], 'region')
        for s in region['subscriptions']:
            if s['subscriptionId'] in comparable:
                close(s['months'][p]['total'], by_id[s['subscriptionId']][p], s['name'])
    pairs = {(r['subscriptionName'], r['serviceName']): r for r in model['subscriptionServiceDrivers']}
    ok_pairs = [p for p in drivers['pairs'] if p.get('status') == 'success']
    for p in ok_pairs:
        c, ref = p['decomposition'], pairs[(p['subscriptionName'], p['serviceName'])]
        close(p['totals']['prev'], ref['prev'])
        close(p['totals']['cur'], ref['cur'])
        close(c['newMeters'] + c['removedMeters'] + c['dailyRateQuantityEffect'] + c['calendarEffect'] + c['unitCostChange'] + c['residual'], p['totals']['delta'])
        close(c['quantityAtPrevUnitCost'], c['calendarEffect'] + c['dailyRateQuantityEffect'])
    ok = [s for s in forecast['subscriptions'] if s['next'].get('fullMonthEstimate') is not None]
    close(math.fsum(s['next']['fullMonthEstimate'] for s in ok), forecast['overall']['next']['fullMonthEstimate'], 'forecast')
    no_cost_without_forecast = all((by_id.get(s['subscriptionId'], {}).get('cur') or 0) == 0 for s in forecast['subscriptions'] if s['next'].get('fullMonthEstimate') is None)
    rs = wb['区域对比']
    total_row = len(region['overall']['byLocation']) + 4
    close(rs.cell(total_row, 2).value, comparison['prev'])
    close(rs.cell(total_row, 3).value, comparison['cur'])
    fs = wb[f'{Ln}费用预测']
    close(fs.cell(len(forecast['subscriptions']) + 4, 5).value, forecast['overall']['next']['fullMonthEstimate'])
    ds = wb['变化原因拆解']
    for i, p in enumerate(ok_pairs, 4):
        close(ds.cell(i, 11).value, ds.cell(i, 5).value)
        close(ds.cell(i, 5).value, p['totals']['delta'])
    chart_rows = list(wb['图表'].iter_rows(min_row=4, max_row=3 + dp + dc, values_only=True))
    close(math.fsum(r[1] or 0 for r in chart_rows), comparison['prev'], '图表上月')
    close(math.fsum(r[2] or 0 for r in chart_rows), comparison['cur'], '图表本月')
    wb.close()
    verification = read_json(fdir / 'verification.json')
    assert verification['totalErrors'] == 0, verification['errors'][:5]
    md = (out / f'{stem(period)}.md').read_text(encoding='utf-8')
    assert f'{comparison["cur"]:,.2f}' in md and f'{comparison["prev"]:,.2f}' in md, 'Markdown 总额与模型不一致'
    for rel in (f'charts/{v}' for v in cost_charts.files(labels).values()):
        assert (out / rel).exists() and f']({rel})' in md, rel
    compact = ''.join(md.split())
    assert '预测代替历史费用' in compact and f'不是{Lc}实际费用减少' in compact
    tokens, tables, active, counts, columns = MarkdownIt('commonmark').enable('table').parse(md), [], False, [], 0
    for t in tokens:
        if t.type == 'table_open':
            active, counts = True, []
        if active and t.type == 'tr_open':
            columns = 0
        if active and t.type in ('td_open', 'th_open'):
            columns += 1
        if active and t.type == 'tr_close':
            counts.append(columns)
        if t.type == 'table_close':
            assert len(set(counts)) == 1, counts
            tables.append(len(counts))
            active = False
    assert len(tables) >= 6, tables
    result = {'status': 'success', 'subscriptions': len(model['subscriptions']), 'comparableSubscriptions': len(comparable),
              'curActualUSD': comparison['cur'], 'prevActualUSD': comparison['prev'], 'deltaUSD': comparison['delta'],
              'formulaCount': verification['formulaCount'], 'recalculatedInExcel': verification.get('recalculatedInExcel'), 'markdownTables': tables,
              'noForecastSubscriptionsHadNoCost': no_cost_without_forecast,
              'checks': ['历史月份不使用预测值', '查询明细与月度、分组合计一致', f'{dp}天/{dc}天日均对比已核对', '读取失败未按0处理',
                         '当前优化建议未写成本月已实现节省', 'Excel 公式结果已核对', '区域合计与订阅合计一致', '变化原因拆解各项合计等于费用变化', '预测合计一致']}
    write_json(fdir / 'audit.json', result)
    log(f'01 核对通过：{len(comparable)} 个订阅，本月 USD {comparison["cur"]:,.2f}，公式 {verification["formulaCount"]} 个，Markdown 表格 {len(tables)} 个')
    return result
