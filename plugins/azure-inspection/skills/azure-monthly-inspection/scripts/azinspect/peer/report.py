"""05 同类资源配置对比报告: Markdown + Excel + 2 charts from findings/peer/baseline.json and narrative/05.json.

Numbers, deviations, adoption and patterns come from the data. Customer-specific explanations (conclusions, top 10,
per-type summaries, template fixes, vendor section, per-resource notes) come from narrative/05.json; without it a
plain default is generated from the data.
"""
from __future__ import annotations

import re
from collections import Counter

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

from ..core import charts as cs
from ..core.excel import verify
from ..core.fmt_md import format_and_check
from ..core.io import read_json, write_json
from ..waf.common import load_scope
from .context import load_narrative

STEM_MD, STEM_XLSX = '05_同类资源配置对比报告', '05_同类资源配置对比明细'
CHARTS = {'method': '05_1_各创建方式高风险偏差比例.png', 'line': '05_2_各业务线高风险偏差比例.png'}
SHORT_TYPE = {'redis': 'Redis', 'mongo': 'MongoDB（vCore）', 'acr': '容器镜像仓库（ACR）', 'ai': 'AI Foundry / OpenAI 账户',
              'appgw': '应用网关', 'vm': '虚拟机', 'vnet': '虚拟网络与子网', 'dbdiag': '数据库与缓存的日志与告警'}
TOP_ORDER = ['aks', 'mysql', 'postgres', 'redis', 'storage', 'appgw', 'vm', 'frontdoor', 'acr', 'mongo', 'vnet']
# Microsoft-recommended target for gaps that no resource of the customer has filled yet (customer independent).
GAP_TARGET = {
    ('aks', 'Entra ID 集成 + 禁用本地账户'): '启用 Entra ID 集成和 Azure RBAC，禁用本地账户',
    ('aks', '网络策略引擎'): '每条业务线确定一种引擎（建议 Cilium），写进模板；Azure NPM 将于 2028-09-30 停止支持',
    ('aks', 'Kubernetes 自动升级'): '至少设为 patch 通道，并配置维护窗口和统一的升级流程',
    ('mysql', '异地冗余备份'): '生产库开启异地冗余备份（只能在创建时选择，结合下次重建或大版本升级处理）',
    ('postgres', '异地冗余备份'): '生产库开启异地冗余备份（只能在创建时选择，结合下次重建或大版本升级处理）',
    ('mysql', 'Azure 告警规则'): '配置 CPU、内存、存储、连接数、可用性告警',
    ('postgres', 'Azure 告警规则'): '配置 CPU、内存、存储、连接数、可用性告警',
    ('dbdiag', 'Azure 告警规则'): '配置 CPU、内存、存储、连接数、可用性告警',
    ('storage', '共享密钥访问'): '关闭共享密钥，改用 Entra ID 授权（先从 IaC 状态账户和备份账户开始）',
    ('keyvault', '清除保护'): '开启清除保护（开启后不能关闭）',
    ('keyvault', '网络访问'): '私有终结点或默认拒绝，并允许受信任的微软服务',
    ('acr', '网络访问'): '升级到 Premium 后加私有终结点或限制来源 IP',
    ('ai', 'API 密钥认证'): '禁用 API 密钥（disableLocalAuth），调用方改用 Entra ID',
    ('ai', '网络访问'): '按需限制网络，承载生产调用的账户优先',
    ('frontdoor', 'WAF 模式'): '观察误报后切换到防护模式',
    ('vm', 'Azure 备份'): '生产机纳入 Azure 备份',
    ('vm', 'Azure Monitor 代理'): '安装 Azure Monitor 代理并配置数据收集规则',
    ('loganalytics', '每日数据上限/数据量告警'): '设置每日上限或数据量告警',
}
PLAIN = [('（lift-and-shift）', ''), ('高影响偏差', '高风险偏差'), ('按基线', '按标准做法'), ('基线', '标准做法')]


def plain(text):
    if text is None:
        return None
    for old, new in PLAIN:
        text = text.replace(old, new)
    return text


def method_label(key):
    return {'未知（无标签、无部署记录）': '既没有标准标签也查不到部署记录的资源', '供应商': '供应商创建的资源',
            '门户向导': '门户向导创建的资源'}.get(key, f'按 {key} 创建的资源')


def cn_date(iso):
    return f'{int(iso[:4])} 年 {int(iso[5:7])} 月 {int(iso[8:10])} 日' if iso and len(iso) >= 10 else '未记录'


def is_high(dev):
    return (dev.get('impact') or '').startswith('高')


def type_name(t):
    return SHORT_TYPE.get(t['typeKey'], t['type'])


def pct(v):
    return int((v or 0) + 0.5)


def ratio(h, n, p):
    return f'{n} 个全部（100%）' if h == n else f'{n} 个中有 {h} 个（{pct(p)}%）'


def table(header, rows):
    out = ['|' + '|'.join(header) + '|', '|' + '|'.join('---' for _ in header) + '|']
    for r in rows:
        assert len(r) == len(header), (header, r)
        out.append('|' + '|'.join(str(c if c is not None else '').replace('|', '／').replace('\n', '；') for c in r) + '|')
    return out


# ---------------------------------------------------------------- per-type figures

def compliance(t):
    """(compliant, total, text) for the per-type table."""
    k, total, b = t['typeKey'], t['population']['total'], t['baseline']
    if k == 'vm':
        pub = sum(int(re.match(r'(\d+)/', d['differingAttributes']['publicIp']['value']).group(1))
                  for d in t['deviations'] if 'publicIp' in d['differingAttributes'])
        m = re.search(r'共 (\d+) 台', t['population'].get('note') or '')
        vms = int(m.group(1)) if m else total
        return 0 if t['deviations'] else total, total, f'{total} 个订阅组中 {total - len(t["deviations"])} 个完全符合（{vms} 台中 {pub} 台直接带公网 IP）'
    if k == 'dbdiag':
        miss = all_ = 0
        for d in t['deviations']:
            a, b2 = map(int, re.match(r'(\d+)/(\d+)', d['differingAttributes']['diag']['value']).groups())
            miss, all_ = miss + a, all_ + b2
        all_ = all_ or total
        return all_ - miss, all_, f'{all_} 个实例中 {all_ - miss} 个导出了日志'
    if not b.get('attributeAdoption'):
        return 0, total, f'{total} 个都不符合（没有内部范例）' if t['uniformGaps'] else f'{total} 个配置一致，无偏差'
    n = len(b['fullyCompliantResources'])
    return n, total, f'{total} 个都不完全符合' if n == 0 else f'{total} 个全部符合' if n == total else f'{total} 个中 {n} 个'


def flat_diff(dev):
    return '\n'.join(f'{a["label"]}：{plain(str(a["value"]))} → {plain(str(a["baseline"]))}' for a in dev['differingAttributes'].values())


def cause_text(dev, extra_notes=(), created_available=True):
    r, parts = dev.get('rootCauseEvidence') or {}, []
    if r.get('creationMethod'):
        parts.append('创建方式：' + r['creationMethod'])
    if r.get('createdTime'):
        parts.append('创建时间：' + r['createdTime'])
    elif not created_available and 'earliestCreated' not in r:
        parts.append('创建时间：未取得')
    if r.get('standardTagged'):
        parts.append(f'带标准创建标签：{r["standardTagged"]} 台')
    if r.get('earliestCreated'):
        parts.append('最早创建：' + r['earliestCreated'])
    pc = r.get('portalChangesAfterCreation') or []
    if pc:
        parts.append(f'创建后在门户上修改 {len(pc)} 次（{pc[0][:10]} 至 {pc[-1][:10]}）')
    parts += list(r.get('notes') or []) + [x for x in extra_notes if x not in (r.get('notes') or [])]
    return plain('\n'.join(parts))


def feas_text(dev):
    out = [dev.get('feasibility') or '']
    fa = dev.get('feasibilityByAttribute') or {}
    if len(fa) > 1:
        out += [f'{dev["differingAttributes"][k]["label"]}：{v}' for k, v in fa.items() if k in dev['differingAttributes']]
    return '\n'.join(x for x in out if x)


def resource_notes(narrative, type_key, name):
    notes = narrative.get('resourceNotes') or {}
    return list(notes.get(f'{type_key}|{name}') or notes.get(name) or [])


# ---------------------------------------------------------------- defaults from data

def default_type_row(t):
    adoption = t['baseline'].get('attributeAdoption') or {}
    standard = '；'.join(f'{a["label"]}：{plain(str(a["baseline"]))}' for a in adoption.values()) or \
        '没有内部范例，按微软推荐：' + '；'.join(f'{k}：{v}' for k, v in (t['baseline'].get('attributes') or {}).items())
    devs = t['deviations']
    if not devs:
        return standard, '没有偏差', '—', '—'
    methods = Counter(((d.get('rootCauseEvidence') or {}).get('creationMethod') or '').split('（')[0] for d in devs)
    exceptions = '、'.join(d['resourceName'] for d in devs[:4]) + (f' 等 {len(devs)} 个' if len(devs) > 4 else '')
    cause = '；'.join(f'{k} {v} 个' for k, v in methods.most_common(3) if k)
    fixes = '；'.join(dict.fromkeys(d.get('feasibility') for d in devs if d.get('feasibility')))
    return standard, exceptions, fixes, cause


def default_top(types, narrative, limit=10):
    order = {k: i for i, k in enumerate(TOP_ORDER)}
    devs = [(t, d) for t in types for d in t['deviations'] if is_high(d)]
    # customer-owned production resources first; vendor-managed ones need the vendor and come last
    devs.sort(key=lambda td: ('供应商' in str(td[1].get('businessLine')), not str(td[1].get('environment')).startswith('生产'),
                              order.get(td[0]['typeKey'], 99), -len(td[1]['differingAttributes'])))
    out = []
    for t, d in devs[:limit]:
        out.append({'resource': f'**{d["resourceName"]}**（{d.get("businessLine")}，{type_name(t)}）', 'difference': flat_diff(d),
                    'cause': cause_text(d, resource_notes(narrative, t['typeKey'], d['resourceName'])),
                    'alignment': f'**{d.get("feasibility")}：**{plain(d.get("recommendation"))}', 'shortTerm': plain(d.get('shortTermMitigation')) or '—'})
    return out


def default_modules(types):
    """Deviations on resources created the standard way point at template defaults."""
    groups = {}
    for t in types:
        for d in t['deviations']:
            if not (d.get('rootCauseEvidence') or {}).get('standardTag'):
                continue
            for k, a in d['differingAttributes'].items():
                g = groups.setdefault((type_name(t), a['label']), {'values': Counter(), 'target': a['baseline'], 'resources': []})
                g['values'][plain(str(a['value']))[:60]] += 1
                g['resources'].append(d['resourceName'])
    return [{'module': f'{tn} 模板', 'current': f'{label}：' + '；'.join(v for v, _ in g['values'].most_common(2)), 'target': plain(str(g['target'])),
             'resources': '、'.join(g['resources'][:6]) + (f' 等 {len(g["resources"])} 个' if len(g['resources']) > 6 else '')}
            for (tn, label), g in sorted(groups.items(), key=lambda kv: -len(kv[1]['resources']))]


def default_conclusions(data, stats, standard_name):
    hi = stats['highImpactDeviations']
    out = []
    bm = hi['byMethod']
    if bm:
        ordered = sorted(bm.items(), key=lambda kv: -(kv[1]['pct'] or 0))
        out.append('**偏差与创建方式有关。**有高风险偏差的资源：' + '，'.join(f'{method_label(k)}{ratio(v["highImpact"], v["resources"], v["pct"])}' for k, v in ordered) + '。')
    bl = {k: v for k, v in hi['byLine'].items() if v['resources'] >= 5}
    if bl:
        ordered = sorted(bl.items(), key=lambda kv: -(kv[1]['pct'] or 0))
        out.append('**各业务线的情况不同。**有高风险偏差的资源：' + '，'.join(f'{k}{ratio(v["highImpact"], v["resources"], v["pct"])}' for k, v in ordered) + '。')
    sv = stats['standardVsOther']
    if sv['standard']['resources']:
        out.append(f'**按 {standard_name} 创建的资源也有偏差。**{sv["standard"]["resources"]} 个中 {sv["standard"]["deviating"]} 个与标准做法不同，'
                   '说明部分模板默认值需要修正，见“需要修正的模板默认值”。')
    return out


# ---------------------------------------------------------------- charts

def bar_chart(rows, title, subtitle, path):
    """rows: [(label, high, total, pct)]; single series, one hue."""
    rows = list(reversed(rows))
    fig, ax = cs.figure(8.6, 0.52 * len(rows) + 1.5)
    cs.style(ax, 'x')
    for i, (label, h, n, p) in enumerate(rows):
        ax.barh(i, p or 0, height=0.55, color=cs.BLUE)
        ax.text((p or 0) + 1.5, i, ratio(h, n, p), va='center', ha='left', fontsize=9, color=cs.INK)
    ax.set_yticks(range(len(rows)), [r[0] for r in rows], fontsize=9.5, color=cs.INK)
    ax.set_xlim(0, 118)
    ax.set_xticks([0, 25, 50, 75, 100], ['0%', '25%', '50%', '75%', '100%'])
    ax.set_title(title, loc='left', fontsize=11.5, color=cs.INK, pad=24)
    ax.text(0, 1.02, subtitle, transform=ax.transAxes, fontsize=8.5, color=cs.INK2, ha='left', va='bottom')
    cs.save(fig, path)


def make_charts(hi, chart_dir):
    m = sorted(((method_label(k).replace('的资源', ''), v['highImpact'], v['resources'], v['pct']) for k, v in hi['byMethod'].items()),
               key=lambda r: -(r[3] or 0))
    total = sum(v['resources'] for v in hi['byMethod'].values())
    out = {}
    if m:
        bar_chart(m, '按创建方式：有高风险偏差的资源占比', f'AKS、数据库、缓存、存储、ACR、应用网关、Front Door，共 {total} 个资源', chart_dir / CHARTS['method'])
        out['method'] = f'charts/{CHARTS["method"]}'
    lines = sorted(((k, v['highImpact'], v['resources'], v['pct']) for k, v in hi['byLine'].items() if v['resources'] >= 5), key=lambda r: -(r[3] or 0))
    if lines:
        bar_chart(lines, '按业务线：有高风险偏差的资源占比', '同一范围，只列资源数不少于 5 个的业务线', chart_dir / CHARTS['line'])
        out['line'] = f'charts/{CHARTS["line"]}'
    return out


# ---------------------------------------------------------------- Markdown

def scope_line(ws, cfg):
    subs = load_scope(ws)
    disabled = [s for s in subs if s.get('state') != 'Enabled']
    return f'{cfg.name} {len(subs) - len(disabled)} 个 Azure 订阅' + (f'（另有 {len(disabled)} 个未启用订阅无法读取）' if disabled else '')


def markdown(ws, cfg, data, narrative, pngs, labels):
    types, stats = data['types'], data['crossTypeStatistics']
    hi, standard_name = stats['highImpactDeviations'], labels['standard']
    n_dev = sum(len(t['deviations']) for t in types)
    total = sum(v['resources'] for v in hi['byMethod'].values())
    total_hi = sum(v['highImpact'] for v in hi['byMethod'].values())
    date = narrative.get('checkDate') or cn_date(data.get('configSnapshot') or str(data.get('analyzedAt') or '')[:10])
    lines = [f'# {cfg.name} Azure 同类资源配置对比报告', '',
             f'**检查时间：{date}。范围：{scope_line(ws, cfg)}。方式：只读，未修改任何资源。**', '',
             f'本报告把同一类资源放在一起比较（共 {len(types)} 类），找出客户自己已经在用的好做法，列出和它不一样的资源，并说明原因和怎么对齐。'
             f'共列出 {n_dev} 项偏差，逐项明细见 Excel“偏差明细”。', '']
    sec = 1
    lines += [f'## {sec}. 结论', ''] + [f'- {x}' for x in (narrative.get('conclusions') or default_conclusions(data, stats, standard_name))] + ['']
    notes = narrative.get('chartNotes') or {}
    for key, title in (('method', '各创建方式高风险偏差比例'), ('line', '各业务线高风险偏差比例')):
        if key in pngs:
            lines += [f'![{title}]({pngs[key]})', '']
            if notes.get(key):
                lines += [notes[key], '']
    sec += 1
    lines += [f'## {sec}. 什么是“标准做法”和“偏差”', '',
              f'**标准做法**是客户自己的资源里已经在用、而且符合微软推荐的做法，多数取自按 {standard_name} 创建的资源；某类资源没有内部范例时，直接用微软推荐做法。'
              '**偏差**是参数和标准做法不一样的资源。', '',
              '**高风险偏差**是直接对应安全风险或已经引发故障的偏差，包括：AKS 用节点子网分配 Pod IP，或 API Server 对公网开放；数据库有公网防火墙规则或没有私有终结点；'
              'Redis 没有私有终结点或用明文协议；业务存储账户对全网开放或允许 TLS 1.0；容器镜像仓库开着管理员账号；应用网关或 Front Door 没有有效的 WAF。', '']
    sec += 1
    top = narrative.get('topDeviations') or default_top(types, narrative)
    lines += [f'## {sec}. 最重要的 {len(top)} 项偏差', '',
              '按影响排序：已经引发故障的、对全网开放的、保存重要数据的排在前面。“原地修改”指在现有资源上改配置，不需要重建。', '']
    lines += table(['#', '资源', '与标准做法的差异', '原因', '怎么对齐', '短期措施'],
                   [(i + 1, x['resource'], x['difference'], x['cause'], x['alignment'], x.get('shortTerm') or '—') for i, x in enumerate(top)])
    sec += 1
    lines += ['', f'## {sec}. 各类资源的标准做法与例外', '',
              '“完全符合标准的数量”指所有比较项都符合标准做法的资源数量。虚拟机按订阅分组比较，数据库日志与告警按业务线分组比较。', '']
    summaries = narrative.get('typeSummaries') or {}
    rows = []
    for t in types:
        std, exc, fix, _cause = default_type_row(t)
        s = summaries.get(t['typeKey']) or {}
        rows.append((type_name(t), s.get('standard') or std, compliance(t)[2], s.get('exceptions') or exc, s.get('alignment') or fix))
    lines += table(['资源类型', '标准做法', '完全符合标准的数量', '主要例外和原因', '对齐建议'], rows)
    modules = narrative.get('moduleFixes') or default_modules(types)
    if modules:
        sec += 1
        lines += ['', f'## {sec}. 需要修正的模板默认值', '',
                  f'下面这些偏差出在按 {standard_name} 创建的资源上，说明模板参数本身需要修正。只改现有资源不改模板，以后新建的资源还会带着同样的问题。', '']
        lines += table(['模板/业务线', '当前默认值', '应改为', '影响的资源'], [(m['module'], m['current'], m['target'], m['resources']) for m in modules])
        if narrative.get('moduleNote'):
            lines += ['', narrative['moduleNote']]
    vendor = stats['vendorManaged']
    if vendor['resources']:
        sec += 1
        lines += ['', f'## {sec}. 供应商创建的资源', '']
        vnote = narrative.get('vendorNote')
        if vnote:
            lines += vnote if isinstance(vnote, list) else [vnote]
        else:
            lines += [f'供应商创建的资源共 {vendor["resources"]} 个，其中 {vendor["deviating"]} 个与标准做法不同。这些资源由供应商的自动化维护，'
                      '**客户不宜直接修改**，否则可能被改回或失去支持；建议与供应商一起确认其模板能否对齐标准做法。']
    gaps = [(t, g) for t in types for g in t['uniformGaps'] if (t['typeKey'], g['attribute']) in GAP_TARGET]
    if gaps:
        sec += 1
        lines += ['', f'## {sec}. 没有内部范例的共同缺口', '',
                  '下列配置所有同类资源都一样缺，没有可以参照的内部资源，按微软推荐统一补齐，并写进模板默认值。', '']
        lines += table(['资源类型', '缺口', '现状', '微软推荐做法'],
                       [(type_name(t), g['attribute'], plain(g.get('note')), GAP_TARGET[(t['typeKey'], g['attribute'])]) for t, g in gaps])
    sec += 1
    lines += ['', f'## {sec}. 方法与局限', '', '**怎么判断原因：**', '',
              f'- 标准创建标签：资源带 `{labels["standardTag"]}` 标签，视为按 {standard_name} 创建。' if labels['standardTag'] else '- 未配置标准创建标签。',
              '- 创建时间：读取每个资源的创建时间，和同类资源的时间先后对比。' if data.get('createdTimeAvailable') else '- 创建时间：本次未取得。',
              '- 门户部署记录：资源组部署历史中出现门户向导的部署名（如 microsoft.aks-*、CreateRedisEnterprise-*、PostgreSqlFlexibleServer_*、CreateVm-*），视为门户向导创建。',
              '- 防火墙规则名称：ClientIPAddress_日期、AllowAll_日期 是门户“添加当前客户端 IP”“允许全部”按钮自动生成的格式，说明规则是在门户上加的。',
              '- 环境（生产、测试）按标签、资源名称和订阅名推断，请客户确认。', '', '**局限：**', '']
    lines += [f'- {plain(x)}' for x in data['coverageGaps'] + list(narrative.get('limitations') or [])]
    lines += [f'- 高风险偏差比例只统计了第 1 节所列 {total} 个资源（其中 {total_hi} 个有高风险偏差），虚拟机、网络、Key Vault、AI 账户未计入。',
              '- “可原地修改”“需重建”等判断依据目前的 Azure 功能，实施前请按微软当前文档再确认一次，并先在测试环境演练。']
    return format_and_check('\n'.join(lines) + '\n')


# ---------------------------------------------------------------- Excel

def write_excel(path, cfg, data, narrative, labels):
    types, stats = data['types'], data['crossTypeStatistics']
    hi = stats['highImpactDeviations']
    expected = {}
    wb = Workbook()
    ov = wb.active
    ov.title = '概览'
    ov.append([f'{cfg.name} 同类资源配置对比明细'])
    ov.append(['“符合数”指所有比较项都符合标准做法的资源数；虚拟机按订阅分组，数据库日志与告警按业务线分组。高风险偏差＝直接对应安全风险或已引发故障的偏差。'])
    ov.append(['类型', '资源数', '标准做法', '符合数', '符合比例', '偏差数', '高风险偏差数', '主要原因', '备注'])
    summaries = narrative.get('typeSummaries') or {}
    for t in types:
        c, n, _txt = compliance(t)
        std, _exc, _fix, cause = default_type_row(t)
        s = summaries.get(t['typeKey']) or {}
        ov.append([type_name(t), n, s.get('standard') or std, c, round(c / n, 3) if n else None, len(t['deviations']),
                   sum(map(is_high, t['deviations'])), s.get('cause') or cause, s.get('note') or t['population'].get('note')])
    last = ov.max_row
    ov.append(['合计', None, None, None, None, f'=SUM(F4:F{last})', f'=SUM(G4:G{last})', None, None])
    expected[f'概览!F{last + 1}'] = sum(len(t['deviations']) for t in types)
    expected[f'概览!G{last + 1}'] = sum(sum(map(is_high, t['deviations'])) for t in types)
    for r in range(4, last + 1):
        ov.cell(r, 5).number_format = '0%'

    dv = wb.create_sheet('偏差明细')
    dv.append(['偏差明细（每行一个资源或一组资源）'])
    dv.append(['“与标准做法的差异”按“项目：当前值 → 标准做法”列出。关联问题编号对应可靠性与安全、卓越运营与性能效率报告中的问题。负责人、客户决定、处理状态、完成日期由客户填写。'])
    dv.append(['编号', '资源类型', '订阅', '资源组', '资源', '环境', '业务线', '风险', '与标准做法的差异', '原因', '关联问题编号', '建议', '可行性', '短期措施',
               '需要客户确认', '负责人', '客户决定', '处理状态', '完成日期'])
    for t in types:
        for i, d in enumerate(t['deviations'], 1):
            dv.append([f'{t["typeKey"].upper()}-{i:02d}', type_name(t), d.get('subscriptionName'), d.get('resourceGroup'), d['resourceName'],
                       d.get('environment'), d.get('businessLine'), '高风险' if is_high(d) else '一般', flat_diff(d),
                       cause_text(d, resource_notes(narrative, t['typeKey'], d['resourceName']), data.get('createdTimeAvailable', True)),
                       '、'.join(d.get('linkedFindingIds') or []),
                       plain(d.get('recommendation')), feas_text(d), plain(d.get('shortTermMitigation')), plain(d.get('needsConfirmation')),
                       None, None, None, None])
    for row in dv.iter_rows(min_row=4):
        if row[7].value == '高风险':
            row[7].fill = PatternFill('solid', fgColor='F8CBAD')

    ad = wb.create_sheet('各属性采用情况')
    ad.append(['各比较项的采用情况'])
    ad.append(['“已采用”指该项已符合标准做法的资源数（虚拟机为订阅组数，数据库日志为业务线组数）。'])
    ad.append(['资源类型', '比较项', '标准做法', '参与比较数', '已采用', '采用比例', '说明'])
    for t in types:
        for a in t['baseline']['attributeAdoption'].values():
            ad.append([type_name(t), a['label'], plain(str(a['baseline'])), a['applicable'], a['followedBy'],
                       (a['followedPct'] or 0) / 100 if a['followedPct'] is not None else None, plain(a.get('note'))])
        if not t['baseline']['attributeAdoption']:
            ad.append([type_name(t), '（无内部范例）', plain('；'.join(f'{k}：{v}' for k, v in t['baseline']['attributes'].items())),
                       t['population']['total'], None, None, plain(t['baseline']['rationale'])])
    for row in ad.iter_rows(min_row=4):
        row[5].number_format = '0%'

    mo = wb.create_sheet('模板修正')
    mo.append(['需要修正的模板默认值'])
    mo.append([plain(narrative.get('moduleNote')) or '按标准方式创建的资源中出现的偏差，说明模板参数需要修正。'])
    mo.append(['模板/业务线', '当前默认值', '应改为', '影响的资源', '负责人', '处理状态'])
    for m in narrative.get('moduleFixes') or default_modules(types):
        mo.append([str(m['module']).replace('`', ''), str(m['current']).replace('`', ''), str(m['target']).replace('`', ''),
                   str(m['resources']).replace('`', ''), None, None])

    gp = wb.create_sheet('共同缺口')
    gp.append(['没有内部范例的共同缺口'])
    gp.append(['所有同类资源都一样缺的配置，按微软推荐统一补齐，并写进模板默认值。'])
    gp.append(['资源类型', '缺口', '符合数', '总数', '说明', '微软推荐做法'])
    for t in types:
        for g in t['uniformGaps']:
            gp.append([type_name(t), g['attribute'], g.get('followedBy'), g.get('of'), plain(g.get('note')), GAP_TARGET.get((t['typeKey'], g['attribute']))])

    pa = wb.create_sheet('规律分析')
    pa.append(['偏差的规律：按创建方式、业务线、区域、创建月份'])
    pa.append([f'高风险偏差统计范围：AKS、MySQL、PostgreSQL、Redis、MongoDB、存储账户、ACR、应用网关、Front Door 共 '
               f'{sum(v["resources"] for v in hi["byMethod"].values())} 个资源；“任一偏差”统计范围另含虚拟网络。'])
    pa.append(['按创建方式', '资源数', '有高风险偏差', '比例'])
    blocks = [('按创建方式', [(k, v['resources'], v['highImpact'], v['pct']) for k, v in hi['byMethod'].items()]),
              ('按业务线', [(k, v['resources'], v['highImpact'], v['pct']) for k, v in hi['byLine'].items()]),
              ('按区域', [(k, v['resources'], v['highImpact'], v['pct']) for k, v in hi['byRegion'].items()]),
              ('按创建月份（任一偏差）', [(k, v['resources'], v['deviating'], v['deviationPct']) for k, v in stats['byCreationMonth'].items()]),
              ('按业务线（任一偏差）', [(k, v['resources'], v['deviating'], v['deviationPct']) for k, v in stats['byBusinessLine'].items()])]
    chart_rows = None
    for bi, (name, rows) in enumerate(blocks):
        if bi:
            pa.append([])
            pa.append([name, '资源数', '有偏差' if '任一' in name else '有高风险偏差', '比例'])
            for c in pa[pa.max_row]:
                c.fill = PatternFill('solid', fgColor='DCE6F1')
        first = pa.max_row + 1
        for r in (rows if '月份' in name else sorted(rows, key=lambda r: -(r[3] or 0))):
            pa.append([r[0], r[1], r[2], (r[3] or 0) / 100 if r[3] is not None else None])
            pa.cell(pa.max_row, 4).number_format = '0%'
        if bi == 0:
            chart_rows = (first, pa.max_row)
    sv = stats['standardVsOther']
    pa.append([])
    pa.append([f'{sv["standardName"]} 与其他方式（任一偏差）', f'{sv["standardName"]} 资源数', '其中有偏差', '其他资源数', '其中有偏差'])
    for c in pa[pa.max_row]:
        c.fill = PatternFill('solid', fgColor='DCE6F1')
    for k, v in sv['byType'].items():
        pa.append([k, v['standard']['n'], v['standard']['deviating'], v['other']['n'], v['other']['deviating']])
    pa.append(['合计', sv['standard']['resources'], sv['standard']['deviating'], sv['other']['resources'], sv['other']['deviating']])
    if chart_rows and chart_rows[1] >= chart_rows[0]:
        pa.auto_filter.ref = f'A3:D{chart_rows[1]}'
        ch = BarChart()
        ch.type, ch.gapWidth, ch.height, ch.width, ch.legend = 'bar', 80, 6, 14, None
        ch.title = '按创建方式：有高风险偏差的比例'
        ch.add_data(Reference(pa, min_col=4, min_row=chart_rows[0], max_row=chart_rows[1]), titles_from_data=False)
        ch.set_categories(Reference(pa, min_col=1, min_row=chart_rows[0], max_row=chart_rows[1]))
        ch.series[0].graphicalProperties.solidFill = ch.series[0].graphicalProperties.line.solidFill = cs.BLUE[1:]
        ch.x_axis.scaling.orientation = 'maxMin'
        ch.x_axis.delete = ch.y_axis.delete = False
        ch.y_axis.number_format = '0%'
        pa.add_chart(ch, 'G3')

    me = wb.create_sheet('方法与局限')
    me.append(['方法与局限'])
    me.append(['怎么找标准做法、怎么判断原因，以及这次比较没有覆盖的内容。'])
    me.append(['类别', '说明'])
    for x in [plain(data['method']['summary']), data['method']['standardDetection'], data['method']['creationMethodDetection'],
              '可行性分类：原地修改；原地修改（需维护窗口或应用配合）；新建节点池/新实例替换（不重建整体）；需重建并迁移；需供应商确认。']:
        me.append(['方法', x])
    for x in data['coverageGaps'] + list(narrative.get('limitations') or []):
        me.append(['局限', plain(x)])
    _style(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return expected


def _style(wb):
    widths = {'概览': {'A': 30, 'C': 60, 'D': 10, 'E': 10, 'H': 40, 'I': 40},
              '偏差明细': {'A': 11, 'B': 18, 'C': 18, 'D': 16, 'E': 26, 'F': 16, 'G': 16, 'H': 8, 'I': 60, 'J': 60, 'K': 22, 'L': 60, 'M': 32, 'N': 40, 'O': 40},
              '各属性采用情况': {'A': 22, 'B': 24, 'C': 44, 'G': 50}, '模板修正': {'A': 26, 'B': 40, 'C': 44, 'D': 60},
              '共同缺口': {'A': 22, 'B': 26, 'E': 60, 'F': 50}, '规律分析': {'A': 34, 'B': 16, 'C': 14, 'D': 18, 'E': 14}, '方法与局限': {'A': 10, 'B': 120}}
    for s in wb:
        s.sheet_view.showGridLines = False
        s.freeze_panes = 'A4'
        if s.title != '规律分析':
            s.auto_filter.ref = f'A3:{L(s.max_column)}{s.max_row}'
        s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=s.max_column)
        s.merge_cells(start_row=2, start_column=1, end_row=2, end_column=s.max_column)
        for row in s:
            for c in row:
                c.font = Font(name='Arial', size=10)
                c.alignment = Alignment(wrap_text=True, vertical='top')
                if c.row == 1:
                    c.fill = PatternFill('solid', fgColor='17365D')
                    c.font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
                elif c.row == 3:
                    c.fill = PatternFill('solid', fgColor='DCE6F1')
                    c.font = Font(name='Arial', size=10, bold=True)
        s.row_dimensions[1].height, s.row_dimensions[2].height = 24, 45
        for i in range(1, s.max_column + 1):
            s.column_dimensions[L(i)].width = 14
        for col, w in widths.get(s.title, {}).items():
            s.column_dimensions[col].width = w


# ---------------------------------------------------------------- build / audit

def build(ws, cfg, narrative=None, log=print):
    fdir = ws.month_dir('findings') / 'peer'
    path = fdir / 'baseline.json'
    if not path.exists():
        log('05：没有同类对比分析结果（findings/peer/baseline.json），跳过')
        return None
    data = read_json(path)
    narrative = narrative if narrative is not None else load_narrative(ws)
    labels = {'standard': data['methodLabels']['standard'], 'standardTag': (data['method'].get('standardDetection') or '').replace('标签 ', '')
              if (data['method'].get('standardDetection') or '').startswith('标签') else None}
    out = ws.month_dir('reports')
    pngs = make_charts(data['crossTypeStatistics']['highImpactDeviations'], out / 'charts')
    xlsx = out / f'{STEM_XLSX}.xlsx'
    expected = write_excel(xlsx, cfg, data, narrative, labels)
    md = markdown(ws, cfg, data, narrative, pngs, labels)
    (out / f'{STEM_MD}.md').write_text(md, encoding='utf-8')
    result = verify(xlsx, expected, fdir / 'verification.json')
    n = sum(len(t['deviations']) for t in data['types'])
    high = sum(sum(map(is_high, t['deviations'])) for t in data['types'])
    meta = {'types': len(data['types']), 'deviations': n, 'highImpact': high, 'formulaCount': result['formulaCount'],
            'recalculatedInExcel': result['recalculatedInExcel']}
    write_json(fdir / 'report.json', meta)
    log(f'05 报告：{len(data["types"])} 类资源，{n} 处偏差（高风险 {high}）；Excel 公式 {result["formulaCount"]} 个'
        + ('，已用 Excel 复算核对' if result['recalculatedInExcel'] else '，本机未能用 Excel 复算'))
    return meta


def audit(ws, log=print):
    from openpyxl import load_workbook
    fdir, out = ws.month_dir('findings') / 'peer', ws.month_dir('reports')
    if not (fdir / 'report.json').exists():
        return None
    data, meta = read_json(fdir / 'baseline.json'), read_json(fdir / 'report.json')
    n = sum(len(t['deviations']) for t in data['types'])
    assert meta['deviations'] == n, '报告记录的偏差数与分析结果不一致'
    wb = load_workbook(out / f'{STEM_XLSX}.xlsx', read_only=True)
    rows = [r for r in wb['偏差明细'].iter_rows(min_row=4, values_only=True) if r[0]]
    assert len(rows) == n, f'Excel 偏差明细 {len(rows)} 行，分析结果 {n} 处'
    assert sum(1 for r in rows if r[7] == '高风险') == meta['highImpact'], '高风险偏差数不一致'
    wb.close()
    md = (out / f'{STEM_MD}.md').read_text(encoding='utf-8')
    assert f'共列出 {n} 项偏差' in md, 'Markdown 偏差总数与分析结果不一致'
    verification = read_json(fdir / 'verification.json')
    assert verification['totalErrors'] == 0, verification['errors'][:5]
    log(f'05 核对通过：{n} 处偏差，Excel 偏差明细逐行对应，公式 {verification["formulaCount"]} 个')
    return meta
