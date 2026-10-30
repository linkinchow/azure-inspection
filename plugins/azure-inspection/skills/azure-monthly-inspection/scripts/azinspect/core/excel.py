"""Excel helpers: shared styling and formula verification.

Every number a customer reads in a workbook is either a value or a formula whose result we know in
advance. `verify` recalculates the workbook in Excel (Windows + pywin32 + Excel installed) and compares
each formula result with the expected value. Without Excel it only checks that every expected formula
is present, and says why in the result (recalculationError); the workbook still recalculates when the customer opens it.
AZINSPECT_NO_EXCEL=1 skips Excel (unit tests only: report audits that read cached values need the recalculation).
"""
from __future__ import annotations

import json
import math
import os
import platform
import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

TITLE_FILL, HEADER_FILL, BAND_FILL = '17365D', 'DCE6F1', 'F5F8FC'
SEVERITY_FILL = {'高': 'F8CBAD', '中': 'FFE699', '低': 'E2EFDA'}
MONEY = '#,##0.00;(#,##0.00);"-"'


def style_sheet(ws, header_row: int = 3, band: bool = False, width: int = 18) -> None:
    """Title row, note row and header row styling used by every report workbook."""
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name='Arial', size=10)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if c.row == 1:
                c.fill = PatternFill('solid', fgColor=TITLE_FILL)
                c.font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
            elif c.row == header_row:
                c.fill = PatternFill('solid', fgColor=HEADER_FILL)
                c.font = Font(name='Arial', size=10, bold=True)
            elif band and c.row > header_row and c.row % 2 == 0:
                c.fill = PatternFill('solid', fgColor=BAND_FILL)
    from openpyxl.utils import get_column_letter
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def recalculation_timeout(size_bytes: int) -> int:
    """Seconds to wait for Excel: 150 for ordinary workbooks, about 30 per MB for large ones, at most an hour."""
    return int(min(3600, max(150, 30 * size_bytes / 1_000_000)))


def _recalculate_with_excel(path: Path) -> tuple[bool, str | None]:
    """(recalculated, reason when not)."""
    if os.environ.get('AZINSPECT_NO_EXCEL'):
        return False, 'AZINSPECT_NO_EXCEL=1，跳过 Excel 复算'
    if platform.system() != 'Windows':
        return False, '不是 Windows，无法调用 Excel'
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False, '未安装 pywin32，无法调用 Excel'
    timeout = recalculation_timeout(path.stat().st_size)
    pythoncom.CoInitialize()
    excel = book = None
    try:
        excel = win32com.client.DispatchEx('Excel.Application')
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3  # disable macros
        book = excel.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=False)
        excel.CalculateFullRebuild()
        deadline = time.monotonic() + timeout
        while excel.CalculationState != 0:
            if time.monotonic() > deadline:
                raise TimeoutError(f'Excel 在 {timeout} 秒内没有完成复算')
            time.sleep(.2)
        book.Save()
        book.Close(SaveChanges=False)
        book = None
        return True, None
    except Exception as e:  # Excel missing, file not opened or COM failure: fall back to the structural check and keep the reason
        return False, f'Excel 复算失败：{type(e).__name__}: {e}'[:500]
    finally:
        # after a disconnect even looking up Close or Quit fails; that must not turn the fallback into a failed build
        try:
            if book is not None:
                book.Close(SaveChanges=False)
        except Exception:  # noqa: BLE001 - the result and its reason are already decided
            pass
        try:
            if excel is not None:
                excel.Quit()
        except Exception:  # noqa: BLE001
            pass
        pythoncom.CoUninitialize()


def unrecalculated_formula(cell, recalculated: bool) -> bool:
    """True when an audit cannot read a formula's result: without Excel recalculation the file holds only the formula text.
    verify() has already checked that each expected formula is present; audits skip the value comparison for such cells."""
    return not recalculated and cell.data_type == 'f'


def verify(path: Path, expected: dict, result_path: Path | None = None) -> dict:
    """expected: {'Sheet!A1': number | {'text': str}}. Raises AssertionError on any mismatch."""
    path = Path(path)
    recalculated, reason = _recalculate_with_excel(path)
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True) if recalculated else None
    errors, count = [], 0
    for ws in formulas:
        for row in ws.iter_rows():
            for c in row:
                key = f'{ws.title}!{c.coordinate}'
                if c.data_type != 'f':
                    continue
                count += 1
                if key not in expected:
                    errors.append({'cell': key, 'error': 'formula without expected value'})
                    continue
                if not recalculated:
                    continue
                cached = values[ws.title][c.coordinate]
                target = expected[key]
                if cached.data_type == 'e':
                    errors.append({'cell': key, 'error': cached.value})
                elif isinstance(target, dict) and 'text' in target:
                    if cached.value != target['text']:
                        errors.append({'cell': key, 'expected': target['text'], 'actual': cached.value})
                elif not isinstance(cached.value, (int, float)) or not math.isclose(cached.value, target, abs_tol=1e-6, rel_tol=1e-9):
                    errors.append({'cell': key, 'expected': target, 'actual': cached.value})
    missing = sorted(set(expected) - {f'{ws.title}!{c.coordinate}' for ws in formulas for row in ws.iter_rows() for c in row if c.data_type == 'f'})
    errors += [{'cell': k, 'error': 'expected formula missing'} for k in missing]
    result = {'status': 'success' if not errors else 'errors_found', 'recalculatedInExcel': recalculated,
              'formulaCount': count, 'totalErrors': len(errors), 'errors': errors[:200], 'workbook': path.name}
    if not recalculated:
        result['recalculationError'] = reason
        result['note'] = '本机无法用 Excel 复算，只检查了公式是否齐全；客户打开文件时 Excel 会自动重算。'
    if result_path:
        Path(result_path).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    assert not errors, errors[:20]
    return result
